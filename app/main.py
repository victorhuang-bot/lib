from fastapi import FastAPI, Request, Response, HTTPException, Form
from fastapi.responses import HTMLResponse, JSONResponse, StreamingResponse, FileResponse, RedirectResponse
from fastapi.staticfiles import StaticFiles
from pathlib import Path
import sqlite3, os, secrets, hashlib, hmac, base64, io, json, asyncio, csv, smtplib, ssl, re
try:
    import psycopg
    from psycopg.rows import dict_row
    from psycopg_pool import ConnectionPool
except Exception:
    psycopg=None
    ConnectionPool=None

from email.message import EmailMessage
from datetime import datetime, timedelta, date, timezone
from zoneinfo import ZoneInfo
import qrcode
from openpyxl import Workbook
from openpyxl.drawing.image import Image as XLImage
from reportlab.pdfgen import canvas
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.cidfonts import UnicodeCIDFont
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.lib.utils import ImageReader

BASE = Path(__file__).resolve().parent.parent
DATA_DIR = Path(os.getenv('DATA_DIR', str(BASE / 'data')))
DB = DATA_DIR / 'app.db'
DATABASE_URL = (os.getenv('DATABASE_URL') or '').strip()
USE_POSTGRES = bool(DATABASE_URL)
APP_VERSION='V18.3.15'

_PREFILL_CACHE = {}
_PREFILL_CACHE_TTL_SECONDS = 45

def _prefill_cache_get(key):
    import time
    item=_PREFILL_CACHE.get(key)
    if not item:
        return None
    ts,value=item
    if time.time()-ts > _PREFILL_CACHE_TTL_SECONDS:
        _PREFILL_CACHE.pop(key,None)
        return None
    return value

def _prefill_cache_set(key,value):
    import time
    _PREFILL_CACHE[key]=(time.time(),value)

def _prefill_cache_clear(service_date=None):
    if service_date is None:
        _PREFILL_CACHE.clear()
        return
    for k in list(_PREFILL_CACHE):
        if isinstance(k,tuple) and service_date in k:
            _PREFILL_CACHE.pop(k,None)

# V18.3.9: Neon connection pool.
# min_size=0 lets Neon scale to zero when truly idle; max_size limits free-tier pressure.
_PG_POOL = None
def pg_pool():
    global _PG_POOL
    if not USE_POSTGRES:
        return None
    if ConnectionPool is None:
        raise RuntimeError('DATABASE_URL 已設定，但 psycopg_pool 未安裝')
    if _PG_POOL is None:
        _PG_POOL=ConnectionPool(
            conninfo=DATABASE_URL,
            min_size=0,
            max_size=5,
            timeout=8,
            max_idle=60,
            kwargs={
                'row_factory': dict_row,
                'connect_timeout': 5
            },
            open=True
        )
    return _PG_POOL
APP_ENV = os.getenv('APP_ENV','development').lower()
APP_BASE_URL = os.getenv('APP_BASE_URL','https://lib.moving-match.com').rstrip('/')
DEMO_RESET_LINKS = os.getenv('DEMO_RESET_LINKS','false').lower() == 'true'
DEMO_ACTIVE_BRANCHES = int(os.getenv('DEMO_ACTIVE_BRANCHES','3'))

ADMIN_USERNAME = os.getenv('ADMIN_USERNAME','moving')
SECRETARY_USERNAME = os.getenv('SECRETARY_USERNAME','lib')
RESET_EMAIL = os.getenv('RESET_EMAIL','lib@moving-match.com')

# V18.3 PostgreSQL Edition official roster
OFFICIAL_ROSTER = {1: [('K14', '李科永'), ('K13S', '士林替'), ('KOB2', '社子島'), ('L32', '關渡宮'), ('L12', '稻香'), ('L23', '秀山'), ('L11', '北投'), ('L14', '清江'), ('L15', '吉利'), ('L21', '永明'), ('L13', '石牌'), ('K12', '天母'), ('K11', '葫蘆堵'), ('F12', '大同'), ('KOB', '百齡智慧'), ('A13', '三民')], 2: [('AFB2', '小巨蛋'), ('DFB', '行天宮'), ('COB', '東區地下街智慧圖書館'), ('A14', '中崙'), ('E12', '城中'), ('C23', '龍安'), ('E11', '王貫英'), ('EOB2', '古亭智慧'), ('E31', '南機場'), ('GOB', '太陽圖書'), ('EFB2', '小南門'), ('F11', '延平'), ('F13', '建成'), ('D12', '長安'), ('F21', '蘭州'), ('AOB', '松山機場')], 3: [('C01', '總館'), ('C01M', '視聽室'), ('B11', '永春'), ('BFB', '臺北市政府'), ('B13', '六合'), ('B12', '三興'), ('EFB', '臺北車站'), ('D11', '中山'), ('D21', '恆安'), ('A12', '民生'), ('A15', '啟明'), ('A11', '松山替'), ('AFB', '松山車站')], 4: [('B14', '廣慈'), ('I21', '龍華'), ('I31', '北原'), ('I11', '南港'), ('IFB', '南港車站'), ('I22', '親子美育'), ('I12', '舊莊'), ('J12', '東湖'), ('J11', '內湖'), ('J13', '西湖'), ('J14', '西中'), ('D13', '大直')], 5: [('H23', '萬芳'), ('H14', '萬興'), ('H12', '木柵'), ('H22', '安康'), ('H16', '力行'), ('H15', '文山'), ('H17', '景新'), ('H11', '景美'), ('C11', '道藩'), ('C22', '成功'), ('CFB', '信義安和'), ('C21', '延吉')], 6: [('G14', '萬華'), ('G13', '西園'), ('G11', '龍山'), ('G21', '柳鄉'), ('G12', '東園')]}
OFFICIAL_BRANCH_CODES = {c for stops in OFFICIAL_ROSTER.values() for c,_ in stops}

# V18.3.2 official default route-driver mapping.
DEFAULT_ROUTE_DRIVER_NAMES = {
    1: '許春芳',
    2: '陳錦隆',
    3: '彭運土',
    4: '林聖原',
    5: '張閔傑',
    6: '陳錦隆',
}

def initial_secret(name, development_default):
    value=os.getenv(name)
    if value:
        return value
    if APP_ENV != 'production':
        return development_default
    raise RuntimeError(f'Missing required production environment variable: {name}')

STATIC = BASE / 'app' / 'static'
app = FastAPI(title='圖書物流配送暨電子簽收管理系統 MVP')
app.mount('/static', StaticFiles(directory=STATIC), name='static')

subscribers = set()

TAIPEI_TZ = ZoneInfo("Asia/Taipei")
def now():
    # All newly written operational timestamps use Taiwan Standard Time (UTC+8).
    return datetime.now(TAIPEI_TZ).isoformat(timespec='seconds')

def future_iso(**kwargs):
    # Expiry timestamps MUST use the same timezone/format as now().
    # Mixing naive server UTC datetime.now() with Asia/Taipei now() caused
    # newly-created driver activation QR codes to be judged expired immediately.
    return (datetime.now(TAIPEI_TZ) + timedelta(**kwargs)).isoformat(timespec='seconds')

def report_time(v):
    """Display report timestamps in Taiwan Standard Time.
    Legacy Render records were stored as naive UTC; aware values keep their timezone.
    """
    if not v: return ''
    try:
        dt=datetime.fromisoformat(str(v).replace('Z','+00:00'))
        if dt.tzinfo is None:
            dt=dt.replace(tzinfo=timezone.utc)
        return dt.astimezone(TAIPEI_TZ).strftime('%Y-%m-%d %H:%M:%S')
    except Exception:
        return str(v)

def today(): return date.today().isoformat()
class _PgCursor:
    def __init__(self, cur):
        self.cur=cur
        self._lastrowid=None
    def fetchone(self):
        r=self.cur.fetchone()
        if r and isinstance(r,dict) and len(r)==1 and 'id' in r:
            self._lastrowid=r['id']
        return r
    def fetchall(self): return self.cur.fetchall()
    @property
    def rowcount(self): return self.cur.rowcount
    @property
    def lastrowid(self):
        if self._lastrowid is None:
            r=self.cur.fetchone()
            if r:
                self._lastrowid=(r.get('id') if isinstance(r,dict) else r[0])
        return self._lastrowid
    def __iter__(self): return iter(self.cur)

class _PgCompat:
    def __init__(self, url):
        self.pool=pg_pool()
        self.con=self.pool.getconn(timeout=8)
        self.con.row_factory=dict_row
        # Neon pooler rejects statement_timeout as a startup parameter.
        # Apply it only after a connection is acquired.
        try:
            with self.con.cursor() as cur:
                cur.execute("SET statement_timeout TO '10s'")
        except Exception:
            # Timeout configuration must never prevent the app from starting.
            try: self.con.rollback()
            except Exception: pass
    def _q(self, sql):
        # Application SQL uses sqlite qmark parameters; psycopg uses %s.
        return sql.replace('?', '%s')
    def execute(self, sql, params=()):
        st=sql.strip()
        if st.upper().startswith('PRAGMA '):
            class Empty:
                rowcount=0
                lastrowid=None
                def fetchone(self): return None
                def fetchall(self): return []
                def __iter__(self): return iter([])
            return Empty()
        if st.upper()=='BEGIN IMMEDIATE':
            st='BEGIN'
        cur=self.con.cursor()
        # Auto-return generated identity id where application asks for lastrowid.
        m=re.match(r'\s*INSERT\s+INTO\s+([A-Za-z_][A-Za-z0-9_]*)',st,re.I)
        identity_tables={'users','password_resets','drivers','driver_devices','driver_activation_tokens',
          'branches','daily_routes','deliveries','corrections','audit_logs','daily_reports',
          'email_recipients','email_logs','delivery_exceptions','global_closures','route_handoffs',
          'delivery_driver_assignments','route_segment_signatures','routes'}
        if m and m.group(1).lower() in identity_tables and ' RETURNING ' not in st.upper():
            st=st.rstrip().rstrip(';')+' RETURNING id'
        cur.execute(self._q(st), params)
        return _PgCursor(cur)
    def executescript(self, script):
        # Convert the SQLite bootstrap DDL to PostgreSQL identity columns.
        script=re.sub(r'\bid INTEGER PRIMARY KEY\b',
                      'id BIGINT GENERATED BY DEFAULT AS IDENTITY PRIMARY KEY',
                      script)
        for stmt in script.split(';'):
            if stmt.strip():
                self.execute(stmt)
    def commit(self): self.con.commit()
    def rollback(self): self.con.rollback()
    def close(self):
        if getattr(self,'con',None) is not None:
            try:
                # Never return an aborted transaction to the pool.
                try:
                    if self.con.info.transaction_status != psycopg.pq.TransactionStatus.IDLE:
                        self.con.rollback()
                except Exception:
                    pass
                self.pool.putconn(self.con)
            finally:
                self.con=None

def db():
    if USE_POSTGRES:
        if psycopg is None or ConnectionPool is None:
            raise RuntimeError('DATABASE_URL 已設定，但 PostgreSQL driver/pool 尚未安裝')
        last=None
        for attempt in range(3):
            try:
                return _PgCompat(DATABASE_URL)
            except Exception as e:
                last=e
                if attempt < 2:
                    import time
                    time.sleep(2)
        raise RuntimeError(f'PostgreSQL 暫時無法連線：{type(last).__name__}: {last}')
    # Local development fallback only.
    DATA_DIR.mkdir(parents=True, exist_ok=True)
    con=sqlite3.connect(DB, timeout=30); con.row_factory=sqlite3.Row
    con.execute('PRAGMA foreign_keys=ON'); con.execute('PRAGMA journal_mode=WAL')
    return con

def public_base_url(req: Request):
    if APP_BASE_URL:
        return APP_BASE_URL
    proto=req.headers.get('x-forwarded-proto') or req.url.scheme
    host=req.headers.get('x-forwarded-host') or req.headers.get('host')
    if host:
        return f'{proto}://{host}'.rstrip('/')
    return str(req.base_url).rstrip('/')

def hash_secret(s, salt=None):
    salt = salt or secrets.token_bytes(16)
    dig = hashlib.pbkdf2_hmac('sha256', s.encode(), salt, 120000)
    return base64.b64encode(salt).decode()+'$'+base64.b64encode(dig).decode()
def verify_secret(s, stored):
    try:
        a,b=stored.split('$'); salt=base64.b64decode(a); exp=base64.b64decode(b)
        got=hashlib.pbkdf2_hmac('sha256', s.encode(), salt, 120000)
        return hmac.compare_digest(exp,got)
    except: return False
def thash(s): return hashlib.sha256(s.encode()).hexdigest()

def smtp_config():
    # V18.2: non-secret Google Workspace defaults are safe fallbacks.
    # Render only needs to hold SMTP_PASSWORD as a secret; every value can still be overridden.
    raw_port=(os.getenv('SMTP_PORT') or '587').strip() or '587'
    try: port=int(raw_port)
    except ValueError: port=587
    user=(os.getenv('SMTP_USER') or 'victor.huang@moving-match.com').strip()
    sender=(os.getenv('SMTP_FROM') or 'lib@moving-match.com').strip()
    return {
        'host': (os.getenv('SMTP_HOST') or 'smtp.gmail.com').strip(),
        'port': port,
        'user': user,
        'password': os.getenv('SMTP_PASSWORD','').strip(),
        'from': sender or user,
        'tls': (os.getenv('SMTP_TLS') or 'true').lower() not in ('0','false','no'),
    }

def smtp_diagnostics():
    cfg=smtp_config()
    # Port/TLS and the non-secret Workspace values have application defaults.
    # Only fields required by the effective runtime config are reported missing.
    checks={
        'SMTP_HOST': bool(cfg['host']),
        'SMTP_PORT': bool(cfg['port']),
        'SMTP_USER': bool(cfg['user']),
        'SMTP_PASSWORD': bool(cfg['password']),
        'SMTP_FROM': bool(cfg['from']),
        'SMTP_TLS': isinstance(cfg['tls'], bool),
    }
    source={
        'SMTP_HOST': 'Render' if os.getenv('SMTP_HOST') else '系統預設',
        'SMTP_PORT': 'Render' if os.getenv('SMTP_PORT') else '系統預設',
        'SMTP_USER': 'Render' if os.getenv('SMTP_USER') else '系統預設',
        'SMTP_PASSWORD': 'Render' if os.getenv('SMTP_PASSWORD') else '未設定',
        'SMTP_FROM': 'Render' if os.getenv('SMTP_FROM') else '系統預設',
        'SMTP_TLS': 'Render' if os.getenv('SMTP_TLS') else '系統預設',
    }
    return {'checks':checks,'source':source,'missing':[k for k,v in checks.items() if not v]}

def smtp_send(subject, body, to_list, cc_list=None, attachments=None):
    cfg=smtp_config(); cc_list=cc_list or []; attachments=attachments or []
    if not cfg['host'] or not cfg['from'] or not cfg['password']:
        raise RuntimeError('SMTP 尚未完整設定：'+', '.join(smtp_diagnostics()['missing']))
    msg=EmailMessage(); msg['Subject']=subject; msg['From']=cfg['from']; msg['To']=', '.join(to_list)
    if cc_list: msg['Cc']=', '.join(cc_list)
    msg.set_content(body)
    for filename,data,mime in attachments:
        maintype,subtype=mime.split('/',1); msg.add_attachment(data,maintype=maintype,subtype=subtype,filename=filename)
    if cfg['port']==465:
        with smtplib.SMTP_SSL(cfg['host'],cfg['port'],context=ssl.create_default_context(),timeout=30) as x:
            if cfg['user']: x.login(cfg['user'],cfg['password'])
            x.send_message(msg)
    else:
        with smtplib.SMTP(cfg['host'],cfg['port'],timeout=30) as x:
            if cfg['tls']: x.starttls(context=ssl.create_default_context())
            if cfg['user']: x.login(cfg['user'],cfg['password'])
            x.send_message(msg)

def init_db():
    DB.parent.mkdir(parents=True, exist_ok=True)
    c=db()
    c.executescript('''
    CREATE TABLE IF NOT EXISTS users(id INTEGER PRIMARY KEY, username TEXT UNIQUE, password_hash TEXT, role TEXT, email TEXT, is_active INTEGER DEFAULT 1, failed_count INTEGER DEFAULT 0, locked_until TEXT);
    CREATE TABLE IF NOT EXISTS sessions(token_hash TEXT PRIMARY KEY, user_id INTEGER, expires_at TEXT);
    CREATE TABLE IF NOT EXISTS password_resets(id INTEGER PRIMARY KEY, user_id INTEGER, token_hash TEXT UNIQUE, expires_at TEXT, used_at TEXT, created_at TEXT);
    CREATE TABLE IF NOT EXISTS routes(id INTEGER PRIMARY KEY, code TEXT UNIQUE, name TEXT, active INTEGER DEFAULT 1);
    CREATE TABLE IF NOT EXISTS drivers(id INTEGER PRIMARY KEY, name TEXT, active INTEGER DEFAULT 1, pin_hash TEXT);
    CREATE TABLE IF NOT EXISTS driver_devices(id INTEGER PRIMARY KEY, driver_id INTEGER, device_key TEXT UNIQUE, activated_at TEXT, revoked_at TEXT, last_seen_at TEXT);
    CREATE TABLE IF NOT EXISTS driver_activation_tokens(id INTEGER PRIMARY KEY, driver_id INTEGER, token_hash TEXT UNIQUE, created_at TEXT, expires_at TEXT, used_at TEXT, revoked_at TEXT, created_by INTEGER);
    CREATE TABLE IF NOT EXISTS branches(id INTEGER PRIMARY KEY, code TEXT UNIQUE, name TEXT, route_id INTEGER, stop_order INTEGER, active INTEGER DEFAULT 1, pin_hash TEXT, access_token_hash TEXT UNIQUE, qr_created_at TEXT, address TEXT DEFAULT '', phone TEXT DEFAULT '', contact_name TEXT DEFAULT '', contact_info TEXT DEFAULT '', delivery_weekdays TEXT DEFAULT '1,2,3,4,5', delivery_frequency TEXT DEFAULT '每週固定');
    CREATE TABLE IF NOT EXISTS daily_routes(id INTEGER PRIMARY KEY, service_date TEXT, route_id INTEGER, driver_id INTEGER, status TEXT DEFAULT 'ACTIVE', driver_signature TEXT, driver_signed_at TEXT, UNIQUE(service_date,route_id));
    CREATE TABLE IF NOT EXISTS deliveries(id INTEGER PRIMARY KEY, service_date TEXT, daily_route_id INTEGER, branch_id INTEGER, status TEXT DEFAULT 'WAITING_SECRETARY', document_original INTEGER, document_final INTEGER, outbound_original INTEGER, outbound_final INTEGER, inbound_final INTEGER, note_final TEXT, signer_name TEXT, branch_signed_at TEXT, branch_signature TEXT, correction_signature TEXT, correction_signer_name TEXT, correction_reason TEXT, corrected_at TEXT, driver_confirmed_at TEXT, row_version INTEGER DEFAULT 1, UNIQUE(service_date,branch_id));
    CREATE TABLE IF NOT EXISTS corrections(id INTEGER PRIMARY KEY, delivery_id INTEGER, requested_by_driver_id INTEGER, requested_at TEXT, fields_json TEXT, driver_note TEXT, status TEXT, resolved_at TEXT);
    CREATE TABLE IF NOT EXISTS audit_logs(id INTEGER PRIMARY KEY, actor_type TEXT, actor_id TEXT, role TEXT, action TEXT, entity_type TEXT, entity_id TEXT, before_json TEXT, after_json TEXT, reason TEXT, created_at TEXT);
    CREATE TABLE IF NOT EXISTS branch_sessions(token_hash TEXT PRIMARY KEY, branch_id INTEGER, delivery_id INTEGER, expires_at TEXT);
    CREATE TABLE IF NOT EXISTS driver_sessions(token_hash TEXT PRIMARY KEY, driver_id INTEGER, device_id INTEGER, expires_at TEXT);
    CREATE TABLE IF NOT EXISTS daily_reports(id INTEGER PRIMARY KEY, service_date TEXT UNIQUE, secretary_signature TEXT, secretary_signed_at TEXT, status TEXT DEFAULT 'OPEN', locked_at TEXT);
    CREATE TABLE IF NOT EXISTS email_recipients(id INTEGER PRIMARY KEY, email TEXT, recipient_type TEXT DEFAULT 'TO', active INTEGER DEFAULT 1, created_at TEXT);
    CREATE TABLE IF NOT EXISTS email_logs(id INTEGER PRIMARY KEY, report_type TEXT, period TEXT, recipients TEXT, status TEXT, error_message TEXT, sent_at TEXT, created_at TEXT);
    CREATE TABLE IF NOT EXISTS app_settings(key TEXT PRIMARY KEY, value TEXT);
    CREATE TABLE IF NOT EXISTS delivery_exceptions(id INTEGER PRIMARY KEY, service_date TEXT, branch_id INTEGER, exception_type TEXT, reason TEXT, created_by INTEGER, created_at TEXT, UNIQUE(service_date,branch_id));
    CREATE TABLE IF NOT EXISTS global_closures(id INTEGER PRIMARY KEY, service_date TEXT UNIQUE, reason TEXT, created_by INTEGER, created_at TEXT);
    CREATE TABLE IF NOT EXISTS route_handoffs(id INTEGER PRIMARY KEY, service_date TEXT, daily_route_id INTEGER, from_driver_id INTEGER, to_driver_id INTEGER, start_delivery_id INTEGER, reason TEXT, note TEXT, created_by INTEGER, created_at TEXT);
    CREATE TABLE IF NOT EXISTS delivery_driver_assignments(id INTEGER PRIMARY KEY, delivery_id INTEGER UNIQUE, driver_id INTEGER, assigned_at TEXT, handoff_id INTEGER);
    CREATE TABLE IF NOT EXISTS route_segment_signatures(id INTEGER PRIMARY KEY, service_date TEXT, daily_route_id INTEGER, driver_id INTEGER, signature TEXT, signed_at TEXT, UNIQUE(service_date,daily_route_id,driver_id));
    ''')
    if not c.execute('SELECT 1 FROM users').fetchone():
        admin_password=initial_secret('ADMIN_INITIAL_PASSWORD','85017306')
        secretary_password=initial_secret('SECRETARY_INITIAL_PASSWORD','03751080')
        branch_pin=initial_secret('DEMO_BRANCH_PIN','1234')
        c.execute('INSERT INTO users(username,password_hash,role,email) VALUES(?,?,?,?)',(ADMIN_USERNAME,hash_secret(admin_password),'ADMIN',RESET_EMAIL))
        c.execute('INSERT INTO users(username,password_hash,role,email) VALUES(?,?,?,?)',(SECRETARY_USERNAME,hash_secret(secretary_password),'SECRETARY',RESET_EMAIL))
        for i in range(1,7): c.execute('INSERT INTO routes(id,code,name) VALUES(?,?,?)',(i,str(i),f'路線{i}'))
        names=['松山分館','民生分館','中山分館','北投分館','石牌分館','大安分館','士林分館','萬華分館','信義分館','內湖分館','南港分館','文山分館']
        for i in range(1,81):
            name=names[i-1] if i<=len(names) else f'示範分館{i:02d}'
            route=((i-1)%6)+1; order=((i-1)//6)+1; tok=secrets.token_urlsafe(24)
            active=1 if i<=DEMO_ACTIVE_BRANCHES else 0
            c.execute('INSERT INTO branches(code,name,route_id,stop_order,active,pin_hash,access_token_hash,qr_created_at) VALUES(?,?,?,?,?,?,?,?)',(f'B{i:03d}',name,route,order,active,hash_secret(branch_pin),thash(tok),now()))
            # raw bootstrap token only for demo retrieval is derived by rotation endpoint; initial QR regenerated below through helper table impossible, so store demo token in audit
            c.execute('INSERT INTO audit_logs(actor_type,role,action,entity_type,entity_id,after_json,created_at) VALUES(?,?,?,?,?,?,?)',('SYSTEM','SYSTEM','SEED_BRANCH_TOKEN','BRANCH',str(i),json.dumps({'token':tok}),now()))
        c.commit()
    migrate_corrections_for_repeat_requests(c)
    migrate_route_secretary_signatures(c)
    migrate_official_routes_and_branches(c)
    migrate_default_route_drivers(c)
    stabilize_accounts_and_today_drivers(c)
    if USE_POSTGRES:
        # Explicit seed IDs 1..6 must advance identity sequences before later inserts.
        for tbl in ('routes','users','drivers','branches'):
            c.execute("SELECT setval(pg_get_serial_sequence(?, 'id'), COALESCE((SELECT MAX(id) FROM "+tbl+"),1), true)",(tbl,))
        c.commit()
    apply_demo_branch_limit(c)
    ensure_today(c)
    reset_test_day_once(c)
    c.close()

def migrate_route_secretary_signatures(c):
    if USE_POSTGRES:
        c.execute('ALTER TABLE daily_routes ADD COLUMN IF NOT EXISTS secretary_signature TEXT')
        c.execute('ALTER TABLE daily_routes ADD COLUMN IF NOT EXISTS secretary_signed_at TEXT')
    else:
        cols={r['name'] for r in c.execute("PRAGMA table_info(daily_routes)").fetchall()}
        if 'secretary_signature' not in cols: c.execute('ALTER TABLE daily_routes ADD COLUMN secretary_signature TEXT')
        if 'secretary_signed_at' not in cols: c.execute('ALTER TABLE daily_routes ADD COLUMN secretary_signed_at TEXT')
    c.commit()

def migrate_corrections_for_repeat_requests(c):
    if USE_POSTGRES:
        return
    row=c.execute("SELECT sql FROM sqlite_master WHERE type='table' AND name='corrections'").fetchone()
    sql=(row['sql'] or '') if row else ''
    if 'delivery_id INTEGER UNIQUE' in sql or 'UNIQUE' in sql.upper():
        c.execute('ALTER TABLE corrections RENAME TO corrections_old')
        c.execute('CREATE TABLE corrections(id INTEGER PRIMARY KEY, delivery_id INTEGER, requested_by_driver_id INTEGER, requested_at TEXT, fields_json TEXT, driver_note TEXT, status TEXT, resolved_at TEXT)')
        c.execute('INSERT INTO corrections(id,delivery_id,requested_by_driver_id,requested_at,fields_json,driver_note,status,resolved_at) SELECT id,delivery_id,requested_by_driver_id,requested_at,fields_json,driver_note,status,resolved_at FROM corrections_old')
        c.execute('DROP TABLE corrections_old')
        c.commit()

def migrate_official_routes_and_branches(c):
    branch_pin=os.getenv('DEMO_BRANCH_PIN') or ('1234' if APP_ENV!='production' else None)
    for rid in range(1,7):
        row=c.execute('SELECT id FROM routes WHERE id=?',(rid,)).fetchone()
        if row: c.execute('UPDATE routes SET code=?,name=?,active=1 WHERE id=?',(str(rid),f'路線{rid}',rid))
        else: c.execute('INSERT INTO routes(id,code,name,active) VALUES(?,?,?,1)',(rid,str(rid),f'路線{rid}'))
    for rid,stops in OFFICIAL_ROSTER.items():
        for order,(code,name) in enumerate(stops,1):
            row=c.execute('SELECT id FROM branches WHERE code=?',(code,)).fetchone()
            if row:
                c.execute('UPDATE branches SET name=?,route_id=?,stop_order=?,active=1 WHERE id=?',(name,rid,order,row['id']))
            else:
                if not branch_pin: raise RuntimeError('Missing required production environment variable: DEMO_BRANCH_PIN')
                raw=secrets.token_urlsafe(24)
                cur=c.execute('INSERT INTO branches(code,name,route_id,stop_order,active,pin_hash,access_token_hash,qr_created_at) VALUES(?,?,?,?,1,?,?,?)',
                              (code,name,rid,order,hash_secret(branch_pin),thash(raw),now()))
                bid=cur.lastrowid
                c.execute('INSERT INTO audit_logs(actor_type,role,action,entity_type,entity_id,after_json,created_at) VALUES(?,?,?,?,?,?,?)',
                          ('SYSTEM','SYSTEM','SEED_OFFICIAL_BRANCH_TOKEN','BRANCH',str(bid),json.dumps({'token':raw},ensure_ascii=False),now()))
    # Old B001-B080 demo branches are disabled, never deleted.
    c.execute("UPDATE branches SET active=0 WHERE code LIKE 'B___' AND length(code)=4")
    c.execute("INSERT INTO app_settings(key,value) VALUES('OFFICIAL_ROSTER_V183_PG',?) ON CONFLICT(key) DO UPDATE SET value=excluded.value",(now(),))
    c.commit()


def migrate_default_route_drivers(c):
    # Keep support/legacy drivers; add the five official drivers if missing.
    driver_ids={}
    for rid,name in DEFAULT_ROUTE_DRIVER_NAMES.items():
        row=c.execute('SELECT id FROM drivers WHERE name=? ORDER BY active DESC,id LIMIT 1',(name,)).fetchone()
        if not row:
            cur=c.execute('INSERT INTO drivers(name,active) VALUES(?,1)',(name,))
            did=cur.lastrowid
        else:
            did=row['id']
            c.execute('UPDATE drivers SET active=1 WHERE id=?',(did,))
        driver_ids[rid]=did
        c.execute(
            "INSERT INTO app_settings(key,value) VALUES(?,?) ON CONFLICT(key) DO UPDATE SET value=excluded.value",
            (f'DEFAULT_ROUTE_DRIVER_{rid}',str(did))
        )

    # Apply defaults to today's route only if that route has not started.
    d=today()
    for rid,did in driver_ids.items():
        dr=c.execute(
            'SELECT id,driver_signed_at FROM daily_routes WHERE service_date=? AND route_id=?',
            (d,rid)
        ).fetchone()
        if not dr:
            continue
        started=c.execute(
            """SELECT 1 FROM deliveries
               WHERE daily_route_id=? AND
                 (document_original IS NOT NULL OR outbound_original IS NOT NULL OR
                  branch_signed_at IS NOT NULL OR driver_confirmed_at IS NOT NULL)
               LIMIT 1""",
            (dr['id'],)
        ).fetchone()
        if not started and not dr['driver_signed_at']:
            c.execute('UPDATE daily_routes SET driver_id=? WHERE id=?',(did,dr['id']))

    c.execute(
        "INSERT INTO app_settings(key,value) VALUES('DEFAULT_ROUTE_DRIVERS_V1832',?) "
        "ON CONFLICT(key) DO UPDATE SET value=excluded.value",
        (now(),)
    )
    c.commit()


def stabilize_accounts_and_today_drivers(c):
    """
    V18.3.4 startup-safe migration:
    - Ensure both ADMIN / SECRETARY accounts exist without overwriting existing passwords.
    - One-time force today's 6 routes to the requested official default drivers.
    - Deactivate old demo/test drivers instead of deleting them, preserving historical references.
    """
    # Core accounts: create only when missing.
    admin=c.execute("SELECT id FROM users WHERE username=?",(ADMIN_USERNAME,)).fetchone()
    if not admin:
        pw=initial_secret('ADMIN_INITIAL_PASSWORD','85017306')
        c.execute(
            'INSERT INTO users(username,password_hash,role,email,is_active) VALUES(?,?,?,?,1)',
            (ADMIN_USERNAME,hash_secret(pw),'ADMIN',RESET_EMAIL)
        )
    sec=c.execute("SELECT id FROM users WHERE username=?",(SECRETARY_USERNAME,)).fetchone()
    if not sec:
        pw=initial_secret('SECRETARY_INITIAL_PASSWORD','03751080')
        c.execute(
            'INSERT INTO users(username,password_hash,role,email,is_active) VALUES(?,?,?,?,1)',
            (SECRETARY_USERNAME,hash_secret(pw),'SECRETARY',RESET_EMAIL)
        )

    # One-time correction for the already-created current day.
    # This intentionally fixes today's old 王/李/陳/林/張 assignment.
    key='FORCE_TODAY_DEFAULT_DRIVERS_V1834'
    if not c.execute('SELECT 1 FROM app_settings WHERE key=?',(key,)).fetchone():
        d=today()
        for rid in range(1,7):
            did=default_driver_id(c,rid)
            dr=c.execute(
                'SELECT id FROM daily_routes WHERE service_date=? AND route_id=?',
                (d,rid)
            ).fetchone()
            if dr:
                c.execute('UPDATE daily_routes SET driver_id=? WHERE id=?',(did,dr['id']))
        c.execute(
            'INSERT INTO app_settings(key,value) VALUES(?,?) ON CONFLICT(key) DO UPDATE SET value=excluded.value',
            (key,now())
        )

    # Keep historical integrity: old demo/test drivers are only deactivated.
    legacy_names=('王先生','李先生','陳先生','林先生','張先生','測試司機')
    placeholders=','.join('?' for _ in legacy_names)
    c.execute(
        f"UPDATE drivers SET active=0 WHERE name IN ({placeholders})",
        legacy_names
    )
    c.commit()

def default_driver_id(c, route_id):
    row=c.execute(
        'SELECT value FROM app_settings WHERE key=?',
        (f'DEFAULT_ROUTE_DRIVER_{route_id}',)
    ).fetchone()
    if row:
        try:
            did=int(row['value'])
            if c.execute('SELECT 1 FROM drivers WHERE id=? AND active=1',(did,)).fetchone():
                return did
        except Exception:
            pass

    name=DEFAULT_ROUTE_DRIVER_NAMES.get(int(route_id))
    if name:
        row=c.execute(
            'SELECT id FROM drivers WHERE name=? AND active=1 ORDER BY id LIMIT 1',
            (name,)
        ).fetchone()
        if row:
            return row['id']

    row=c.execute('SELECT id FROM drivers WHERE active=1 ORDER BY id LIMIT 1').fetchone()
    if not row:
        raise RuntimeError('No active driver available')
    return row['id']

def apply_demo_branch_limit(c):
    if c.execute("SELECT 1 FROM app_settings WHERE key='OFFICIAL_ROSTER_V183_PG'").fetchone(): return
    # Apply the 3-branch demo preset once only. After that, admin enable/disable choices are respected.
    key='DEMO_BRANCH_LIMIT_INITIALIZED_V5'
    if c.execute('SELECT 1 FROM app_settings WHERE key=?',(key,)).fetchone():
        return
    n=max(1, DEMO_ACTIVE_BRANCHES)
    c.execute('UPDATE branches SET active=CASE WHEN id<=? THEN 1 ELSE 0 END',(n,))
    c.execute("DELETE FROM branch_sessions WHERE delivery_id IN (SELECT d.id FROM deliveries d JOIN branches b ON b.id=d.branch_id WHERE d.service_date=? AND b.active=0)",(today(),))
    c.execute("DELETE FROM corrections WHERE delivery_id IN (SELECT d.id FROM deliveries d JOIN branches b ON b.id=d.branch_id WHERE d.service_date=? AND b.active=0)",(today(),))
    c.execute("DELETE FROM deliveries WHERE service_date=? AND branch_id IN (SELECT id FROM branches WHERE active=0)",(today(),))
    c.execute('INSERT INTO app_settings(key,value) VALUES(?,?)',(key,now()))
    c.commit()

def reset_test_day_once(c):
    key='V4_THREE_BRANCH_TEST_RESET'
    if c.execute('SELECT 1 FROM app_settings WHERE key=?',(key,)).fetchone():
        return
    # One-time clean restart for the hosted test: reset only today's three active test branches.
    active_ids=[r['id'] for r in c.execute('SELECT id FROM branches WHERE active=1 ORDER BY id LIMIT ?',(max(1,DEMO_ACTIVE_BRANCHES),)).fetchall()]
    if active_ids:
        q=','.join('?' for _ in active_ids)
        dids=[r['id'] for r in c.execute(f'SELECT id FROM deliveries WHERE service_date=? AND branch_id IN ({q})',(today(),*active_ids)).fetchall()]
        if dids:
            qd=','.join('?' for _ in dids)
            c.execute(f'DELETE FROM branch_sessions WHERE delivery_id IN ({qd})',dids)
            c.execute(f'DELETE FROM corrections WHERE delivery_id IN ({qd})',dids)
            c.execute(f"UPDATE deliveries SET status='WAITING_SECRETARY',document_original=NULL,document_final=NULL,outbound_original=NULL,outbound_final=NULL,inbound_final=NULL,note_final=NULL,signer_name=NULL,branch_signed_at=NULL,branch_signature=NULL,correction_signature=NULL,correction_signer_name=NULL,correction_reason=NULL,corrected_at=NULL,driver_confirmed_at=NULL,row_version=row_version+1 WHERE id IN ({qd})",dids)
    c.execute('INSERT INTO app_settings(key,value) VALUES(?,?)',(key,now()))
    c.commit()

def is_report_locked(c, service_date=None):
    d=service_date or today()
    r=c.execute("SELECT status FROM daily_reports WHERE service_date=?",(d,)).fetchone()
    return bool(r and r['status']=='LOCKED')

def branch_expected_on(c,b,service_date):
    if c.execute('SELECT 1 FROM global_closures WHERE service_date=?',(service_date,)).fetchone():
        return False
    ex=c.execute('SELECT exception_type FROM delivery_exceptions WHERE service_date=? AND branch_id=?',(service_date,b['id'])).fetchone()
    if ex:
        return ex['exception_type']=='ADD'
    weekday=str(date.fromisoformat(service_date).isoweekday())
    days=[x.strip() for x in (b['delivery_weekdays'] or '').split(',') if x.strip()]
    return bool(b['active']) and weekday in days

def rebuild_service_date(c, service_date):
    # Idempotent / concurrency-safe day builder for PostgreSQL.
    # It is safe for dashboard, prefill and schedule refresh to call this repeatedly.
    try:
        date.fromisoformat(service_date)
    except:
        raise ValueError('service_date must be YYYY-MM-DD')

    route_map={}
    routes=c.execute('SELECT * FROM routes WHERE active=1 ORDER BY id').fetchall()
    for r in routes:
        dr=c.execute(
            'SELECT * FROM daily_routes WHERE service_date=? AND route_id=?',
            (service_date,r['id'])
        ).fetchone()
        if not dr:
            driver=default_driver_id(c,r['id'])
            if USE_POSTGRES:
                c.execute(
                    '''INSERT INTO daily_routes(service_date,route_id,driver_id)
                       VALUES(?,?,?)
                       ON CONFLICT(service_date,route_id) DO NOTHING''',
                    (service_date,r['id'],driver)
                )
            else:
                try:
                    c.execute(
                        'INSERT INTO daily_routes(service_date,route_id,driver_id) VALUES(?,?,?)',
                        (service_date,r['id'],driver)
                    )
                except Exception:
                    pass
            dr=c.execute(
                'SELECT * FROM daily_routes WHERE service_date=? AND route_id=?',
                (service_date,r['id'])
            ).fetchone()
        if dr:
            route_map[r['id']]=dr['id']

    branches=c.execute('SELECT * FROM branches ORDER BY id').fetchall()
    for b in branches:
        expected=branch_expected_on(c,b,service_date)
        existing=c.execute(
            'SELECT * FROM deliveries WHERE service_date=? AND branch_id=?',
            (service_date,b['id'])
        ).fetchone()

        if expected and not existing and b['route_id'] in route_map:
            if USE_POSTGRES:
                c.execute(
                    '''INSERT INTO deliveries(service_date,daily_route_id,branch_id,status)
                       VALUES(?,?,?,?)
                       ON CONFLICT(service_date,branch_id) DO NOTHING''',
                    (service_date,route_map[b['route_id']],b['id'],'WAITING_SECRETARY')
                )
            else:
                try:
                    c.execute(
                        'INSERT INTO deliveries(service_date,daily_route_id,branch_id,status) VALUES(?,?,?,?)',
                        (service_date,route_map[b['route_id']],b['id'],'WAITING_SECRETARY')
                    )
                except Exception:
                    pass

        elif not expected and existing:
            untouched=(
                existing['document_original'] is None and
                existing['outbound_original'] is None and
                existing['branch_signed_at'] is None and
                existing['driver_confirmed_at'] is None
            )
            if untouched:
                c.execute('DELETE FROM branch_sessions WHERE delivery_id=?',(existing['id'],))
                c.execute('DELETE FROM corrections WHERE delivery_id=?',(existing['id'],))
                c.execute('DELETE FROM deliveries WHERE id=?',(existing['id'],))
    c.commit()

def next_service_date():
    return (date.fromisoformat(today()) + timedelta(days=1)).isoformat()

def ensure_today(c):
    rebuild_service_date(c,today())

print(f'APP_VERSION={APP_VERSION}',flush=True)
init_db()

def audit(c, actor_type, actor_id, role, action, etype, eid, before=None, after=None, reason=None):
    c.execute('INSERT INTO audit_logs(actor_type,actor_id,role,action,entity_type,entity_id,before_json,after_json,reason,created_at) VALUES(?,?,?,?,?,?,?,?,?,?)',(actor_type,str(actor_id or ''),role,action,etype,str(eid),json.dumps(before,ensure_ascii=False) if before is not None else None,json.dumps(after,ensure_ascii=False) if after is not None else None,reason,now()))

def current_user(req):
    auth=req.headers.get('Authorization','')
    tok=auth[7:] if auth.startswith('Bearer ') else req.cookies.get('session')
    if not tok: return None
    c=db(); r=c.execute('SELECT u.* FROM sessions s JOIN users u ON u.id=s.user_id WHERE s.token_hash=? AND s.expires_at>?',(thash(tok),now())).fetchone(); c.close(); return dict(r) if r else None

def require_user(req, roles=None):
    u=current_user(req)
    if not u: raise HTTPException(401,'UNAUTHENTICATED')
    if roles and u['role'] not in roles: raise HTTPException(403,'FORBIDDEN')
    return u

async def publish(event):
    dead=[]
    for q in list(subscribers):
        try:q.put_nowait(event)
        except:dead.append(q)
    for q in dead: subscribers.discard(q)

@app.get('/health')
def health(): return {'ok':True,'environment':APP_ENV}

@app.middleware('http')
async def no_cache_versioned_assets(request:Request, call_next):
    response=await call_next(request)
    if request.url.path=='/' or request.url.path.startswith('/static/') or request.url.path.startswith('/api/secretary/documents/prefill-v3'):
        response.headers['Cache-Control']='no-store, no-cache, must-revalidate, max-age=0'
        response.headers['X-App-Version']=APP_VERSION
    return response

@app.get('/api/version')
def api_version():
    return {'version':APP_VERSION,'database':'postgresql' if USE_POSTGRES else 'sqlite','prefill_api':'v3'}

@app.get('/', response_class=HTMLResponse)
def home():
    return HTMLResponse((STATIC/'index.html').read_text(encoding='utf-8'),headers={'Cache-Control':'no-store, no-cache, must-revalidate, max-age=0'})
@app.get('/branch/{token}', response_class=HTMLResponse)
def branch_page(token:str): return (STATIC/'branch.html').read_text(encoding='utf-8').replace('__BRANCH_TOKEN__',token)
@app.get('/activate-driver/{token}', response_class=HTMLResponse)
def driver_activate_page(token:str): return (STATIC/'activate.html').read_text(encoding='utf-8').replace('__ACTIVATION_TOKEN__',token)
@app.get('/driver', response_class=HTMLResponse)
def driver_page(): return (STATIC/'driver.html').read_text(encoding='utf-8')
@app.get('/reset-password', response_class=HTMLResponse)
def reset_page(): return (STATIC/'reset.html').read_text(encoding='utf-8')

def _login_db(username,password):
    c=db()
    try:
        u=c.execute('SELECT * FROM users WHERE username=? AND is_active=1',(username,)).fetchone()
        if not u or not verify_secret(password,u['password_hash']):
            raise HTTPException(401,'帳號或密碼錯誤')
        tok=secrets.token_urlsafe(32)
        c.execute('INSERT INTO sessions(token_hash,user_id,expires_at) VALUES(?,?,?)',(thash(tok),u['id'],future_iso(hours=12)))
        audit(c,'USER',u['id'],u['role'],'LOGIN','USER',u['id'])
        c.commit()
        return {'ok':True,'role':u['role'],'token':tok}
    except:
        try: c.rollback()
        except: pass
        raise
    finally:
        c.close()

@app.post('/api/auth/login')
async def login(req:Request):
    p=await req.json()
    result=await asyncio.to_thread(_login_db,p.get('username',''),p.get('password',''))
    r=JSONResponse(result)
    r.set_cookie('session',result['token'],httponly=True,samesite='lax',secure=(APP_ENV=='production'),max_age=43200)
    return r
@app.post('/api/auth/logout')
def logout(req:Request):
    auth=req.headers.get('Authorization',''); tok=auth[7:] if auth.startswith('Bearer ') else req.cookies.get('session'); c=db();
    if tok:c.execute('DELETE FROM sessions WHERE token_hash=?',(thash(tok),));c.commit()
    c.close(); r=JSONResponse({'ok':True});r.delete_cookie('session');return r
@app.get('/api/auth/me')
def me(req:Request):
    u=current_user(req); return {'user': {'username':u['username'],'role':u['role']} if u else None}
@app.post('/api/auth/password/forgot')
async def forgot(req:Request):
    p=await req.json(); username=(p.get('username') or '').strip(); c=db(); u=c.execute('SELECT * FROM users WHERE username=?',(username,)).fetchone(); demo=None
    if u:
        raw=secrets.token_urlsafe(32); c.execute('INSERT INTO password_resets(user_id,token_hash,expires_at,created_at) VALUES(?,?,?,?)',(u['id'],thash(raw),future_iso(minutes=30),now())); c.commit()
        reset_url=APP_BASE_URL+f'/reset-password?token={raw}'
        if APP_ENV!='production' or DEMO_RESET_LINKS: demo=reset_url
        try:
            smtp_send('圖書物流系統密碼重設',f'帳號：{username}\n\n請於 30 分鐘內使用以下連結重設密碼：\n{reset_url}\n\n若非本人操作，請忽略此信。',[RESET_EMAIL])
        except Exception as e:
            c.close(); raise HTTPException(500,'密碼重設信寄送失敗：'+str(e))
    c.close(); return {'ok':True,'message':'若帳號存在，系統已寄送重設連結。','demo_reset_url':demo}
@app.post('/api/auth/password/reset')
async def reset(req:Request):
    p=await req.json(); c=db(); r=c.execute('SELECT * FROM password_resets WHERE token_hash=? AND used_at IS NULL AND expires_at>?',(thash(p.get('token','')),now())).fetchone()
    if not r: c.close(); raise HTTPException(400,'重設連結無效或已過期')
    c.execute('UPDATE users SET password_hash=? WHERE id=?',(hash_secret(p.get('new_password','')),r['user_id'])); c.execute('UPDATE password_resets SET used_at=? WHERE id=?',(now(),r['id'])); c.execute('DELETE FROM sessions WHERE user_id=?',(r['user_id'],)); c.commit(); c.close(); return {'ok':True}


@app.post('/api/account/change-password')
async def change_own_password(req:Request):
    u=require_user(req,['ADMIN','SECRETARY']); p=await req.json(); current=p.get('current_password',''); newp=p.get('new_password',''); confirm=p.get('confirm_password','')
    if newp!=confirm: raise HTTPException(400,'兩次新密碼不一致')
    if len(newp)<8: raise HTTPException(400,'新密碼至少 8 碼')
    c=db(); fresh=c.execute('SELECT * FROM users WHERE id=?',(u['id'],)).fetchone()
    if not fresh or not verify_secret(current,fresh['password_hash']): c.close(); raise HTTPException(400,'目前密碼錯誤')
    c.execute('UPDATE users SET password_hash=? WHERE id=?',(hash_secret(newp),u['id'])); c.execute('DELETE FROM sessions WHERE user_id=? AND token_hash!=?',(u['id'], thash((req.headers.get('Authorization','')[7:] if req.headers.get('Authorization','').startswith('Bearer ') else req.cookies.get('session','')))))
    audit(c,'USER',u['id'],u['role'],'CHANGE_OWN_PASSWORD','USER',u['id']); c.commit(); c.close(); return {'ok':True}

@app.post('/api/account/reset-secretary-password')
async def reset_secretary_password(req:Request):
    u=require_user(req,['ADMIN']); p=await req.json(); admin_password=p.get('admin_password',''); newp=p.get('new_password',''); confirm=p.get('confirm_password','')
    if newp!=confirm: raise HTTPException(400,'兩次新密碼不一致')
    if len(newp)<8: raise HTTPException(400,'新密碼至少 8 碼')
    c=db(); admin=c.execute('SELECT * FROM users WHERE id=?',(u['id'],)).fetchone()
    if not admin or not verify_secret(admin_password,admin['password_hash']): c.close(); raise HTTPException(400,'管理者密碼驗證失敗')
    sec=c.execute("SELECT * FROM users WHERE role='SECRETARY' ORDER BY id LIMIT 1").fetchone()
    if not sec: c.close(); raise HTTPException(404,'找不到秘書帳號')
    c.execute('UPDATE users SET password_hash=? WHERE id=?',(hash_secret(newp),sec['id'])); c.execute('DELETE FROM sessions WHERE user_id=?',(sec['id'],)); audit(c,'USER',u['id'],u['role'],'RESET_SECRETARY_PASSWORD','USER',sec['id']); c.commit(); c.close(); return {'ok':True}

@app.get('/api/account/security')
def account_security(req:Request):
    u=require_user(req,['ADMIN','SECRETARY']); c=db(); rows=[dict(x) for x in c.execute("SELECT username,role,email,is_active FROM users WHERE role IN ('ADMIN','SECRETARY') ORDER BY id").fetchall()]; c.close(); return {'current':{'username':u['username'],'role':u['role']},'accounts':rows}

@app.get('/api/dashboard/today')
def dashboard(req:Request):
    require_user(req,['ADMIN','SECRETARY']); c=db(); ensure_today(c)
    rows=c.execute('SELECT status,COUNT(*) n FROM deliveries WHERE service_date=? GROUP BY status',(today(),)).fetchall(); counts={r['status']:r['n'] for r in rows}; total=sum(counts.values()); completed=counts.get('STOP_COMPLETED',0)
    routes=c.execute('''SELECT dr.id,r.code,r.name,d.name driver,
        (SELECT COUNT(*) FROM deliveries x WHERE x.daily_route_id=dr.id) total,
        (SELECT COUNT(*) FROM deliveries x WHERE x.daily_route_id=dr.id AND x.status='STOP_COMPLETED') completed,
        dr.status
        FROM daily_routes dr
        JOIN routes r ON r.id=dr.route_id
        LEFT JOIN drivers d ON d.id=dr.driver_id
        WHERE dr.service_date=?
        ORDER BY CAST(r.code AS INTEGER),r.code''',(today(),)).fetchall()
    c.close(); return {'date':today(),'total':total,'completed':completed,'waiting_branch':counts.get('WAITING_BRANCH',0),'waiting_driver':counts.get('WAITING_DRIVER',0),'waiting_driver_confirm':counts.get('WAITING_DRIVER_CONFIRM',0)+counts.get('WAITING_DRIVER_RECONFIRM',0),'waiting_correction':counts.get('WAITING_BRANCH_CORRECTION',0),'routes':[dict(x) for x in routes]}
@app.get('/api/dashboard/deliveries')
def dashboard_deliveries(req:Request, status:str|None=None, route:str|None=None, search:str|None=None):
    require_user(req,['ADMIN','SECRETARY']); c=db(); q='''SELECT x.*,b.code branch_code,b.name branch_name,r.code route_code,d.name driver_name FROM deliveries x JOIN branches b ON b.id=x.branch_id JOIN daily_routes dr ON dr.id=x.daily_route_id JOIN routes r ON r.id=dr.route_id JOIN drivers d ON d.id=dr.driver_id WHERE x.service_date=?'''; args=[today()]
    if status:q+=' AND x.status=?';args.append(status)
    if route:q+=' AND r.code=?';args.append(route)
    if search:q+=' AND b.name LIKE ?';args.append('%'+search+'%')
    q+=' ORDER BY r.code,b.stop_order'; rows=[dict(x) for x in c.execute(q,args).fetchall()];c.close();return rows

@app.get('/api/branches')
def branches(req:Request):
    require_user(req,['ADMIN','SECRETARY']); c=db(); rows=[dict(x) for x in c.execute("SELECT b.id,b.code,b.name,r.code route_code,b.route_id,b.stop_order,b.active,b.qr_created_at,b.address,b.phone,b.contact_name,b.contact_info,b.delivery_weekdays,b.delivery_frequency,CASE WHEN b.pin_hash IS NULL THEN '未設定' ELSE '已設定' END pin_status FROM branches b JOIN routes r ON r.id=b.route_id ORDER BY b.active DESC,r.code,b.stop_order").fetchall()];c.close();return rows

def raw_branch_token(c,bid):
    # demo-only retrieval of current token from last rotate/seed audit record
    rows=c.execute("SELECT after_json FROM audit_logs WHERE entity_type='BRANCH' AND entity_id=? AND action IN ('SEED_BRANCH_TOKEN','ROTATE_BRANCH_QR') ORDER BY id DESC LIMIT 1",(str(bid),)).fetchone()
    return json.loads(rows['after_json'])['token'] if rows else None
@app.get('/api/branches/{bid}/qr')
def branch_qr(req:Request,bid:int):
    require_user(req,['ADMIN','SECRETARY']); c=db(); b=c.execute('SELECT * FROM branches WHERE id=?',(bid,)).fetchone(); raw=raw_branch_token(c,bid); c.close()
    if not b or not raw: raise HTTPException(404)
    return {'branch':b['name'],'url':public_base_url(req)+'/branch/'+raw}
@app.get('/api/branches/{bid}/qr.png')
def branch_qr_png(req:Request,bid:int):
    require_user(req,['ADMIN','SECRETARY']); c=db(); raw=raw_branch_token(c,bid); c.close();
    if not raw: raise HTTPException(404)
    im=qrcode.make(public_base_url(req)+'/branch/'+raw); bio=io.BytesIO(); im.save(bio,'PNG'); return Response(bio.getvalue(),media_type='image/png')
@app.post('/api/branches/{bid}/qr/rotate')
async def rotate_qr(req:Request,bid:int):
    u=require_user(req,['ADMIN','SECRETARY']); c=db(); raw=secrets.token_urlsafe(24); c.execute('UPDATE branches SET access_token_hash=?,qr_created_at=? WHERE id=?',(thash(raw),now(),bid)); audit(c,'USER',u['id'],u['role'],'ROTATE_BRANCH_QR','BRANCH',bid,after={'token':raw});c.commit();c.close(); await publish({'type':'branch.qr_rotated','branch_id':bid});return {'ok':True}
@app.post('/api/branches/{bid}/pin')
async def set_branch_pin(req:Request,bid:int):
    u=require_user(req,['ADMIN','SECRETARY']); p=await req.json(); pin=p.get('pin','');
    if len(pin)!=4 or not pin.isdigit():raise HTTPException(400,'PIN需為4位數字')
    c=db();c.execute('UPDATE branches SET pin_hash=? WHERE id=?',(hash_secret(pin),bid));audit(c,'USER',u['id'],u['role'],'RESET_BRANCH_PIN','BRANCH',bid);c.commit();c.close();return {'ok':True}


@app.get('/api/qr.png')
def generic_qr(req:Request, data:str):
    require_user(req,['ADMIN','SECRETARY']); im=qrcode.make(data); bio=io.BytesIO(); im.save(bio,'PNG'); return Response(bio.getvalue(),media_type='image/png')

@app.get('/api/drivers')
def drivers(req:Request):
    require_user(req,['ADMIN','SECRETARY']);c=db();rows=[dict(x) for x in c.execute("SELECT id,name,active,CASE WHEN pin_hash IS NULL THEN 0 ELSE 1 END pin_set FROM drivers ORDER BY active DESC,id").fetchall()];c.close();return rows
@app.post('/api/drivers/{did}/activation')
async def activation(req:Request,did:int):
    u=require_user(req,['ADMIN']); raw=secrets.token_urlsafe(32); c=db()
    # Only the newest activation QR remains valid. Older unused QR codes are revoked immediately.
    c.execute('UPDATE driver_activation_tokens SET revoked_at=? WHERE driver_id=? AND used_at IS NULL AND revoked_at IS NULL',(now(),did))
    c.execute('INSERT INTO driver_activation_tokens(driver_id,token_hash,created_at,expires_at,created_by) VALUES(?,?,?,?,?)',(did,thash(raw),now(),future_iso(minutes=10),u['id']))
    audit(c,'USER',u['id'],u['role'],'CREATE_DRIVER_ACTIVATION_QR','DRIVER',did,after={'expires_in_seconds':600})
    c.commit();c.close();return {'url':public_base_url(req)+'/activate-driver/'+raw,'expires_in_seconds':600}
@app.get('/api/driver-activation/{token}')
def activation_info(token:str):
    c=db(); r=c.execute('''SELECT t.*,d.name FROM driver_activation_tokens t JOIN drivers d ON d.id=t.driver_id WHERE t.token_hash=?''',(thash(token),)).fetchone();c.close();
    if not r or r['used_at'] or r['revoked_at'] or r['expires_at']<=now(): raise HTTPException(410,'此啟用QR Code已失效')
    return {'driver':r['name'],'expires_at':r['expires_at']}
@app.post('/api/driver-activation/{token}')
async def activate_driver(token:str, req:Request):
    p=await req.json(); pin=p.get('pin','');
    if len(pin)!=4 or not pin.isdigit():raise HTTPException(400,'PIN需為4位數字')
    c=db(); c.execute('BEGIN IMMEDIATE'); r=c.execute('SELECT * FROM driver_activation_tokens WHERE token_hash=?',(thash(token),)).fetchone()
    if not r or r['used_at'] or r['revoked_at'] or r['expires_at']<=now(): c.rollback();c.close();raise HTTPException(410,'此啟用QR Code已失效')
    device_key=secrets.token_urlsafe(24); c.execute('UPDATE drivers SET pin_hash=? WHERE id=?',(hash_secret(pin),r['driver_id'])); c.execute('INSERT INTO driver_devices(driver_id,device_key,activated_at,last_seen_at) VALUES(?,?,?,?)',(r['driver_id'],device_key,now(),now())); c.execute('UPDATE driver_activation_tokens SET used_at=? WHERE id=?',(now(),r['id']));c.commit();c.close();return {'ok':True,'device_key':device_key}
@app.post('/api/driver/unlock')
async def driver_unlock(req:Request):
    p=await req.json(); c=db(); dev=c.execute('''SELECT dd.*,d.pin_hash FROM driver_devices dd JOIN drivers d ON d.id=dd.driver_id WHERE dd.device_key=? AND dd.revoked_at IS NULL''',(p.get('device_key',''),)).fetchone()
    if not dev or not verify_secret(p.get('pin',''),dev['pin_hash'] or ''):c.close();raise HTTPException(401,'裝置或PIN錯誤')
    raw=secrets.token_urlsafe(32);c.execute('INSERT INTO driver_sessions(token_hash,driver_id,device_id,expires_at) VALUES(?,?,?,?)',(thash(raw),dev['driver_id'],dev['id'],future_iso(hours=12)));c.commit();c.close();return {'token':raw}
def driver_auth(req):
    auth=req.headers.get('Authorization',''); raw=auth[7:] if auth.startswith('Bearer ') else ''
    c=db();r=c.execute('SELECT * FROM driver_sessions WHERE token_hash=? AND expires_at>?',(thash(raw),now())).fetchone();c.close();
    if not r:raise HTTPException(401,'DRIVER_UNAUTHENTICATED')
    return dict(r)
@app.get('/api/driver/today')
def driver_today(req:Request):
    s=driver_auth(req); c=db(); rows=[dict(x) for x in c.execute('''SELECT x.*,b.name branch_name,b.code branch_code,r.code route_code FROM deliveries x JOIN branches b ON b.id=x.branch_id JOIN daily_routes dr ON dr.id=x.daily_route_id JOIN routes r ON r.id=dr.route_id LEFT JOIN delivery_driver_assignments dda ON dda.delivery_id=x.id WHERE x.service_date=? AND COALESCE(dda.driver_id,dr.driver_id)=? AND b.active=1 ORDER BY r.code,b.stop_order''',(today(),s['driver_id'])).fetchall()];c.close();return rows
@app.patch('/api/driver/deliveries/{did}/outbound')
async def outbound(did:int,req:Request):
    s=driver_auth(req);p=await req.json();qty=int(p.get('qty',0));c=db();
    if is_report_locked(c): c.close(); raise HTTPException(409,'今日日報已鎖定')
    x=c.execute('''SELECT x.*,COALESCE(dda.driver_id,dr.driver_id) driver_id FROM deliveries x JOIN daily_routes dr ON dr.id=x.daily_route_id LEFT JOIN delivery_driver_assignments dda ON dda.delivery_id=x.id WHERE x.id=?''',(did,)).fetchone()
    if not x or x['driver_id']!=s['driver_id']:c.close();raise HTTPException(403)
    if x['status'] not in ('WAITING_DRIVER','WAITING_SECRETARY','WAITING_BRANCH'):c.close();raise HTTPException(409,'分館簽收後不可再修改送出數量')
    status='WAITING_BRANCH' if x['document_original'] is not None else 'WAITING_SECRETARY'; c.execute('UPDATE deliveries SET outbound_original=?,outbound_final=?,status=?,row_version=row_version+1 WHERE id=?',(qty,qty,status,did));audit(c,'DRIVER',s['driver_id'],'DRIVER','SET_OR_EDIT_OUTBOUND','DELIVERY',did,before={'qty':x['outbound_final']},after={'qty':qty});c.commit();c.close();await publish({'type':'delivery.updated','id':did});return {'ok':True}
@app.post('/api/driver/deliveries/{did}/confirm')
async def confirm(did:int,req:Request):
    s=driver_auth(req);c=db();
    if is_report_locked(c): c.close(); raise HTTPException(409,'今日日報已鎖定')
    x=c.execute('''SELECT x.*,COALESCE(dda.driver_id,dr.driver_id) driver_id FROM deliveries x JOIN daily_routes dr ON dr.id=x.daily_route_id LEFT JOIN delivery_driver_assignments dda ON dda.delivery_id=x.id WHERE x.id=?''',(did,)).fetchone()
    if not x or x['driver_id']!=s['driver_id']:c.close();raise HTTPException(403)
    if x['status'] not in ('WAITING_DRIVER_CONFIRM','WAITING_DRIVER_RECONFIRM'):c.close();raise HTTPException(409,'目前狀態不可完成本站')
    c.execute("UPDATE deliveries SET status='STOP_COMPLETED',driver_confirmed_at=?,row_version=row_version+1 WHERE id=?",(now(),did));audit(c,'DRIVER',s['driver_id'],'DRIVER','CONFIRM_STOP','DELIVERY',did);c.commit();c.close();await publish({'type':'delivery.updated','id':did});return {'ok':True}
@app.post('/api/driver/deliveries/{did}/request-correction')
async def request_correction(did:int,req:Request):
    s=driver_auth(req); p=await req.json(); c=db();
    if is_report_locked(c): c.close(); raise HTTPException(409,'今日日報已鎖定')
    x=c.execute('''SELECT x.*,COALESCE(dda.driver_id,dr.driver_id) driver_id FROM deliveries x JOIN daily_routes dr ON dr.id=x.daily_route_id LEFT JOIN delivery_driver_assignments dda ON dda.delivery_id=x.id WHERE x.id=?''',(did,)).fetchone()
    if not x or x['driver_id']!=s['driver_id']:c.close();raise HTTPException(403)
    if x['status'] not in ('WAITING_DRIVER_CONFIRM','WAITING_DRIVER_RECONFIRM'):c.close();raise HTTPException(409,'目前狀態不可要求分館更正')
    if c.execute("SELECT 1 FROM corrections WHERE delivery_id=? AND status='PENDING'",(did,)).fetchone():c.close();raise HTTPException(409,'此筆已有待分館處理的更正要求')
    fields=['document','outbound','inbound']
    c.execute('INSERT INTO corrections(delivery_id,requested_by_driver_id,requested_at,fields_json,driver_note,status) VALUES(?,?,?,?,?,?)',(did,s['driver_id'],now(),json.dumps(fields),p.get('note',''),'PENDING'));audit(c,'DRIVER',s['driver_id'],'DRIVER','REQUEST_CORRECTION','DELIVERY',did,after={'fields':fields,'note':p.get('note','')});c.execute("UPDATE deliveries SET status='WAITING_BRANCH_CORRECTION',row_version=row_version+1 WHERE id=?",(did,));c.commit();c.close();await publish({'type':'delivery.updated','id':did});return {'ok':True}
@app.get('/api/driver/routes/today')
def driver_routes_today(req:Request):
    s=driver_auth(req); c=db(); rows=[dict(x) for x in c.execute("""SELECT dr.id,r.code,r.name,dr.status,dr.driver_signed_at,dr.driver_signature,dr.secretary_signature,dr.secretary_signed_at,(SELECT COUNT(*) FROM deliveries x WHERE x.daily_route_id=dr.id) total,(SELECT COUNT(*) FROM deliveries x WHERE x.daily_route_id=dr.id AND x.status='STOP_COMPLETED') completed FROM daily_routes dr JOIN routes r ON r.id=dr.route_id WHERE dr.service_date=? AND dr.driver_id=? AND EXISTS(SELECT 1 FROM deliveries x WHERE x.daily_route_id=dr.id) ORDER BY r.code""",(today(),s['driver_id'])).fetchall()]; c.close(); return rows

@app.get('/api/driver/routes/{rid}/summary')
def driver_route_summary(rid:int,req:Request):
    s=driver_auth(req); c=db()
    dr=c.execute("""SELECT dr.id,dr.service_date,dr.status,dr.driver_signed_at,dr.driver_signature,dr.secretary_signature,dr.secretary_signed_at,dr.secretary_signature,dr.secretary_signed_at,r.code,r.name,d.name driver_name FROM daily_routes dr JOIN routes r ON r.id=dr.route_id JOIN drivers d ON d.id=dr.driver_id WHERE dr.id=? AND dr.driver_id=?""",(rid,s['driver_id'])).fetchone()
    if not dr: c.close(); raise HTTPException(403)
    rows=[dict(x) for x in c.execute("""SELECT x.id,b.code branch_code,b.name branch_name,x.document_final,x.outbound_final,x.inbound_final,x.note_final,x.signer_name,x.correction_signer_name,x.branch_signed_at,x.corrected_at,x.status FROM deliveries x JOIN branches b ON b.id=x.branch_id WHERE x.daily_route_id=? ORDER BY b.stop_order""",(rid,)).fetchall()]
    totals={'document':sum((x['document_final'] or 0) for x in rows),'outbound':sum((x['outbound_final'] or 0) for x in rows),'inbound':sum((x['inbound_final'] or 0) for x in rows)}
    c.close(); return {'route':dict(dr),'stops':rows,'totals':totals}

@app.post('/api/driver/routes/{rid}/sign')
async def sign_route(rid:int,req:Request):
    s=driver_auth(req); p=await req.json();c=db();dr=c.execute('SELECT * FROM daily_routes WHERE id=? AND driver_id=?',(rid,s['driver_id'])).fetchone();
    if not dr:c.close();raise HTTPException(403)
    n=c.execute("SELECT COUNT(*) n FROM deliveries WHERE daily_route_id=? AND status!='STOP_COMPLETED'",(rid,)).fetchone()['n']
    if n:c.close();raise HTTPException(409,'尚有站點未完成')
    if not p.get('signature',''): c.close(); raise HTTPException(400,'司機路線簽名必填')
    if is_report_locked(c,dr['service_date']): c.close(); raise HTTPException(409,'日報已鎖定')
    c.execute("UPDATE daily_routes SET status='DRIVER_SIGNED',driver_signature=?,driver_signed_at=? WHERE id=?",(p.get('signature',''),now(),rid));audit(c,'DRIVER',s['driver_id'],'DRIVER','DRIVER_ROUTE_SIGN','DAILY_ROUTE',rid,after={'signed_at':now()});c.commit();c.close();await publish({'type':'route.updated','id':rid});return {'ok':True}

@app.post('/api/branch-session/verify')
async def branch_verify(req:Request):
    p=await req.json(); c=db(); b=c.execute('SELECT * FROM branches WHERE access_token_hash=? AND active=1',(thash(p.get('token','')),)).fetchone()
    if not b or not verify_secret(p.get('pin',''),b['pin_hash']):c.close();raise HTTPException(401,'QR或驗證碼錯誤')
    x=c.execute('SELECT * FROM deliveries WHERE service_date=? AND branch_id=?',(today(),b['id'])).fetchone()
    if not x:c.close();raise HTTPException(404,'今日沒有配送任務')
    raw=secrets.token_urlsafe(32);c.execute('INSERT INTO branch_sessions(token_hash,branch_id,delivery_id,expires_at) VALUES(?,?,?,?)',(thash(raw),b['id'],x['id'],future_iso(minutes=30)));c.commit();c.close();return {'session':raw,'branch':b['name']}
def branch_auth(req):
    auth=req.headers.get('Authorization',''); raw=auth[7:] if auth.startswith('Bearer ') else ''
    c=db();r=c.execute('SELECT * FROM branch_sessions WHERE token_hash=? AND expires_at>?',(thash(raw),now())).fetchone();c.close();
    if not r:raise HTTPException(401,'BRANCH_SESSION_EXPIRED')
    return dict(r)
@app.get('/api/branch-session/today')
def branch_today(req:Request):
    s=branch_auth(req);c=db();x=c.execute('''SELECT x.*,b.name branch_name FROM deliveries x JOIN branches b ON b.id=x.branch_id WHERE x.id=? AND x.branch_id=?''',(s['delivery_id'],s['branch_id'])).fetchone();corr=c.execute("SELECT * FROM corrections WHERE delivery_id=? ORDER BY CASE WHEN status='PENDING' THEN 0 ELSE 1 END,id DESC LIMIT 1",(x['id'],)).fetchone();c.close();out=dict(x);out['correction']=dict(corr) if corr else None;return out
@app.post('/api/branch-session/sign')
async def branch_sign(req:Request):
    s=branch_auth(req);p=await req.json();c=db();
    if is_report_locked(c): c.close(); raise HTTPException(409,'今日日報已鎖定')
    x=c.execute('SELECT * FROM deliveries WHERE id=?',(s['delivery_id'],)).fetchone()
    if x['status']!='WAITING_BRANCH':c.close();raise HTTPException(409,'目前不可第一次簽收')
    vals=(int(p['document']),int(p['outbound']),int(p['inbound']),p.get('note',''),p.get('signer','').strip(),p.get('signature',''))
    if not vals[4] or not vals[5]:c.close();raise HTTPException(400,'姓名與簽名必填')
    c.execute("UPDATE deliveries SET document_final=?,outbound_final=?,inbound_final=?,note_final=?,signer_name=?,branch_signature=?,branch_signed_at=?,status='WAITING_DRIVER_CONFIRM',row_version=row_version+1 WHERE id=?",(*vals,now(),x['id']));audit(c,'BRANCH',s['branch_id'],'BRANCH','BRANCH_SIGN','DELIVERY',x['id'],after={'document':vals[0],'outbound':vals[1],'inbound':vals[2],'note':vals[3],'signer':vals[4]});c.commit();c.close();await publish({'type':'delivery.updated','id':x['id']});return {'ok':True}
@app.post('/api/branch-session/correct')
async def branch_correct(req:Request):
    s=branch_auth(req);p=await req.json();c=db();
    if is_report_locked(c): c.close(); raise HTTPException(409,'今日日報已鎖定')
    x=c.execute('SELECT * FROM deliveries WHERE id=?',(s['delivery_id'],)).fetchone();corr=c.execute("SELECT * FROM corrections WHERE delivery_id=? AND status='PENDING' ORDER BY id DESC LIMIT 1",(x['id'],)).fetchone()
    if x['status']!='WAITING_BRANCH_CORRECTION' or not corr or corr['status']!='PENDING':c.close();raise HTTPException(409,'目前沒有可更正資料')
    if not p.get('reason','').strip() or not p.get('signer','').strip() or not p.get('signature',''):c.close();raise HTTPException(400,'更正原因、姓名與新簽名必填')
    def optional_qty(key,current):
        value=p.get(key,None)
        if value is None or str(value).strip()=='':
            return current
        try:
            value=int(value)
        except (TypeError,ValueError):
            c.close(); raise HTTPException(400,f'{key} 數量格式錯誤')
        if value < 0:
            c.close(); raise HTTPException(400,'數量不可小於 0')
        return value
    document=optional_qty('document',x['document_final'])
    outbound_qty=optional_qty('outbound',x['outbound_final'])
    inbound=optional_qty('inbound',x['inbound_final'])
    before={'document':x['document_final'],'outbound':x['outbound_final'],'inbound':x['inbound_final']}
    c.execute("UPDATE deliveries SET document_final=?,outbound_final=?,inbound_final=?,correction_reason=?,correction_signer_name=?,correction_signature=?,corrected_at=?,status='WAITING_DRIVER_RECONFIRM',row_version=row_version+1 WHERE id=?",(document,outbound_qty,inbound,p['reason'],p['signer'],p['signature'],now(),x['id']))
    c.execute("UPDATE corrections SET status='RESOLVED',resolved_at=? WHERE id=?",(now(),corr['id']))
    audit(c,'BRANCH',s['branch_id'],'BRANCH','BRANCH_CORRECT','DELIVERY',x['id'],before=before,after={'document':document,'outbound':outbound_qty,'inbound':inbound,'reason':p['reason'],'signer':p['signer']})
    c.commit();c.close();await publish({'type':'delivery.updated','id':x['id']});return {'ok':True}


def ensure_prefill_delivery(c, service_date, branch_id):
    b=c.execute('SELECT * FROM branches WHERE id=?',(branch_id,)).fetchone()
    if not b or not branch_expected_on(c,b,service_date):
        raise HTTPException(409,'此分館不是所選日期的配送站點')
    dr=c.execute('SELECT * FROM daily_routes WHERE service_date=? AND route_id=?',(service_date,b['route_id'])).fetchone()
    if not dr:
        driver=default_driver_id(c,b['route_id'])
        c.execute('INSERT INTO daily_routes(service_date,route_id,driver_id) VALUES(?,?,?) ON CONFLICT(service_date,route_id) DO NOTHING',(service_date,b['route_id'],driver))
        dr=c.execute('SELECT * FROM daily_routes WHERE service_date=? AND route_id=?',(service_date,b['route_id'])).fetchone()
    x=c.execute('SELECT * FROM deliveries WHERE service_date=? AND branch_id=?',(service_date,branch_id)).fetchone()
    if not x:
        c.execute("INSERT INTO deliveries(service_date,daily_route_id,branch_id,status) VALUES(?,?,?,'WAITING_SECRETARY') ON CONFLICT(service_date,branch_id) DO NOTHING",(service_date,dr['id'],branch_id))
        x=c.execute('SELECT * FROM deliveries WHERE service_date=? AND branch_id=?',(service_date,branch_id)).fetchone()
    return x

@app.get('/api/secretary/documents/prefill-v3')
def document_prefill_v3(req:Request, service_date:str|None=None):
    require_user(req,['SECRETARY'])
    d=service_date or next_service_date()
    try:
        selected=date.fromisoformat(d)
    except:
        raise HTTPException(400,'日期格式錯誤')
    if selected < date.fromisoformat(today()):
        raise HTTPException(400,'公文預填只能選擇今天或未來日期')

    cache_key=('prefill',d)
    cached=_prefill_cache_get(cache_key)
    if cached is not None:
        return cached

    c=db()
    try:
        closure=c.execute('SELECT reason FROM global_closures WHERE service_date=?',(d,)).fetchone()
        if closure:
            result={'app_version':APP_VERSION,'service_date':d,'is_global_closure':True,'closure_reason':closure['reason'],'count':0,'rows':[]}
            c.close(); _prefill_cache_set(cache_key,result); return result

        branches=c.execute("SELECT b.id,b.code,b.name,b.route_id,b.stop_order,b.delivery_weekdays,b.delivery_frequency,r.code route_code,r.name route_name FROM branches b JOIN routes r ON r.id=b.route_id WHERE b.active=1 AND r.active=1 ORDER BY CAST(r.code AS INTEGER),r.code,b.stop_order").fetchall()
        exceptions=c.execute("SELECT branch_id,exception_type,service_date FROM delivery_exceptions WHERE service_date=?",(d,)).fetchall()
        stop_ids={x['branch_id'] for x in exceptions if x['exception_type']=='STOP'}
        add_ids={x['branch_id'] for x in exceptions if x['exception_type']=='ADD'}

        existing_rows=c.execute("SELECT x.id,x.branch_id,x.status,x.document_final,x.outbound_original,dr.route_id,dr.driver_id,dv.name driver_name FROM deliveries x JOIN daily_routes dr ON dr.id=x.daily_route_id LEFT JOIN drivers dv ON dv.id=dr.driver_id WHERE x.service_date=?",(d,)).fetchall()
        existing_by_branch={x['branch_id']:x for x in existing_rows}

        drivers=c.execute('SELECT id,name FROM drivers WHERE active=1').fetchall()
        driver_name_by_id={x['id']:x['name'] for x in drivers}
        settings=c.execute("SELECT key,value FROM app_settings WHERE key LIKE ?",('DEFAULT_ROUTE_DRIVER_%',)).fetchall()
        default_driver_by_route={}
        for x in settings:
            try:
                default_driver_by_route[int(str(x['key']).split('_')[-1])]=int(x['value'])
            except:
                pass

        selected_weekday=selected.isoweekday()
        rows=[]
        for b in branches:
            bid=b['id']
            fixed=False
            raw=(b['delivery_weekdays'] or '').strip()
            if raw:
                try:
                    fixed=selected_weekday in {int(v.strip()) for v in raw.split(',') if v.strip()}
                except:
                    fixed=False
            expected=(fixed and bid not in stop_ids) or (bid in add_ids)
            if not expected:
                continue

            ex=existing_by_branch.get(bid)
            if ex:
                did,status,doc,outbound=ex['id'],ex['status'],ex['document_final'],ex['outbound_original']
                driver_id,driver_name=ex['driver_id'],ex['driver_name']
            else:
                did,status,doc,outbound=None,'WAITING_SECRETARY',None,None
                driver_id=default_driver_by_route.get(b['route_id'])
                if not driver_id:
                    default_name=DEFAULT_ROUTE_DRIVER_NAMES.get(int(b['route_id']))
                    driver_id=next((x['id'] for x in drivers if x['name']==default_name),None)
                driver_name=driver_name_by_id.get(driver_id) if driver_id else None

            rows.append({'id':did,'branch_id':bid,'service_date':d,'status':status,'document_final':doc,'outbound_original':outbound,'branch_code':b['code'],'branch_name':b['name'],'stop_order':b['stop_order'],'route_code':b['route_code'],'route_name':b['route_name'],'driver_id':driver_id,'driver_name':driver_name})

        result={'app_version':APP_VERSION,'service_date':d,'is_global_closure':False,'closure_reason':None,'count':len(rows),'rows':rows}
        c.close(); _prefill_cache_set(cache_key,result); return result
    except Exception as e:
        try: c.rollback()
        except: pass
        c.close(); print('PREFILL_ERROR',d,type(e).__name__,str(e),flush=True)
        raise HTTPException(500,f'公文預填讀取失敗：{type(e).__name__}: {e}')

@app.post('/api/secretary/documents/prefill-save')
async def document_prefill_save(req:Request):
    u=require_user(req,['SECRETARY']); p=await req.json(); d=str(p.get('service_date') or '')
    try: date.fromisoformat(d); bid=int(p.get('branch_id')); qty=int(p.get('qty'))
    except: raise HTTPException(400,'日期、分館或公文數量格式錯誤')
    if qty<0: raise HTTPException(400,'公文數量不可小於 0')
    return await asyncio.to_thread(_prefill_save_db,u['id'],u['role'],d,bid,qty)

def _prefill_batch_db(user_id,user_role,d,items):
    c=db(); changed=0; total=0
    try:
        for it in items:
            bid=int(it.get('branch_id')); qty=int(it.get('qty'))
            if qty<0: raise HTTPException(400,'公文數量不可小於 0')
            x=ensure_prefill_delivery(c,d,bid)
            if x['outbound_original'] is not None: continue
            c.execute('UPDATE deliveries SET document_original=COALESCE(document_original,?),document_final=?,row_version=row_version+1 WHERE id=?',(qty,qty,x['id']))
            changed+=1; total+=qty
        audit(c,'USER',user_id,user_role,'PREFILL_DOCUMENT_BATCH','DELIVERY',d,after={'service_date':d,'changed':changed,'total':total})
        c.commit()
        _prefill_cache_clear(d)
        return {'ok':True,'changed':changed,'total':total}
    except HTTPException:
        c.rollback(); raise
    except Exception as e:
        c.rollback()
        print('PREFILL_BATCH_ERROR',d,type(e).__name__,str(e),flush=True)
        raise HTTPException(500,f'公文批次預填失敗：{type(e).__name__}: {e}')
    finally:
        c.close()

@app.post('/api/secretary/documents/prefill-batch')
async def document_prefill_batch(req:Request):
    u=require_user(req,['SECRETARY']); p=await req.json(); d=str(p.get('service_date') or ''); items=p.get('items') or []
    try: date.fromisoformat(d)
    except: raise HTTPException(400,'日期格式錯誤')
    return await asyncio.to_thread(_prefill_batch_db,u['id'],u['role'],d,items)

@app.patch('/api/secretary/deliveries/{did}/document')
async def set_doc(did:int,req:Request):
    u=require_user(req,['SECRETARY']); p=await req.json()
    try: qty=int(p.get('qty'))
    except: raise HTTPException(400,'公文數量必須是整數')
    if qty < 0: raise HTTPException(400,'公文數量不可小於 0')
    c=db(); x=c.execute('SELECT * FROM deliveries WHERE id=?',(did,)).fetchone()
    if not x: c.close(); raise HTTPException(404,'找不到配送資料')
    if is_report_locked(c,x['service_date']):
        c.close(); raise HTTPException(409,f"{x['service_date']} 日報已鎖定")
    if x['outbound_original'] is not None:
        c.close(); raise HTTPException(409,'司機已輸入圖書送出數量，公文數量已鎖定，無法再修改')
    if x['status'] not in ('WAITING_SECRETARY','WAITING_DRIVER'):
        c.close(); raise HTTPException(409,'目前配送狀態不可修改公文數量')
    before={'qty':x['document_final']}
    c.execute("UPDATE deliveries SET document_original=?,document_final=?,status='WAITING_DRIVER',row_version=row_version+1 WHERE id=?",(qty,qty,did))
    audit(c,'USER',u['id'],u['role'],'SET_DOCUMENT','DELIVERY',did,before=before,after={'qty':qty})
    c.commit(); c.close(); await publish({'type':'delivery.updated','id':did}); return {'ok':True,'qty':qty,'locked':False}

@app.post('/api/secretary/documents/zero-all')
async def zero_all(req:Request, service_date:str|None=None):
    u=require_user(req,['SECRETARY'])
    d=service_date or today()
    try: date.fromisoformat(d)
    except: raise HTTPException(400,'日期格式錯誤')
    c=db()
    if d>today():
        rebuild_service_date(c,d)
    if is_report_locked(c,d): c.close(); raise HTTPException(409,f'{d} 日報已鎖定')
    cur=c.execute("UPDATE deliveries SET document_original=0,document_final=0,status='WAITING_DRIVER',row_version=row_version+1 WHERE service_date=? AND outbound_original IS NULL AND status IN ('WAITING_SECRETARY','WAITING_DRIVER')",(d,))
    changed=cur.rowcount
    audit(c,'USER',u['id'],u['role'],'ZERO_ALL_DOCUMENTS','DAY',d,after={'changed':changed,'service_date':d})
    c.commit(); c.close()
    if d==today(): await publish({'type':'dashboard.refresh'})
    return {'ok':True,'changed':changed,'service_date':d}

@app.post('/api/secretary/documents/batch')
async def batch_documents(req:Request):
    u=require_user(req,['SECRETARY']); p=await req.json(); items=p.get('items') or []
    if not isinstance(items,list) or not items: raise HTTPException(400,'沒有可存檔的公文數量')
    d=p.get('service_date') or today()
    try: date.fromisoformat(d)
    except: raise HTTPException(400,'日期格式錯誤')
    c=db()
    if d>today():
        rebuild_service_date(c,d)
    if is_report_locked(c,d): c.close(); raise HTTPException(409,f'{d} 日報已鎖定')
    prepared=[]
    for item in items:
        try: did=int(item.get('id')); qty=int(item.get('qty'))
        except: c.close(); raise HTTPException(400,'公文數量格式錯誤')
        if qty<0: c.close(); raise HTTPException(400,'公文數量需為0以上整數')
        x=c.execute('SELECT * FROM deliveries WHERE id=? AND service_date=?',(did,d)).fetchone()
        if not x: c.close(); raise HTTPException(404,f'找不到 {d} 配送資料')
        if x['outbound_original'] is not None: c.close(); raise HTTPException(409,f"配送 #{did} 司機已輸入送出數量，公文已鎖定")
        prepared.append((x,qty))
    for x,qty in prepared:
        before={'qty':x['document_final']}
        c.execute("UPDATE deliveries SET document_original=?,document_final=?,status='WAITING_DRIVER',row_version=row_version+1 WHERE id=?",(qty,qty,x['id']))
        audit(c,'USER',u['id'],u['role'],'BATCH_SET_DOCUMENT','DELIVERY',x['id'],before=before,after={'qty':qty})
    c.commit(); c.close()
    if d==today(): await publish({'type':'delivery.updated','batch':True})
    return {'ok':True,'changed':len(prepared),'total':sum(q for _,q in prepared),'service_date':d}

@app.get('/api/schedule-exceptions')
def schedule_exceptions(req:Request, month:str=''):
    require_user(req,['ADMIN','SECRETARY']); month=month or today()[:7]; c=db()
    rows=[dict(x) for x in c.execute("""SELECT e.id,e.service_date,e.branch_id,e.exception_type,e.reason,e.created_at,b.code branch_code,b.name branch_name,r.code route_code FROM delivery_exceptions e JOIN branches b ON b.id=e.branch_id LEFT JOIN routes r ON r.id=b.route_id WHERE e.service_date LIKE ? ORDER BY e.service_date,r.code,b.stop_order""",(month+'%',)).fetchall()]
    for g in c.execute("SELECT id,service_date,reason,created_at FROM global_closures WHERE service_date LIKE ? ORDER BY service_date",(month+'%',)).fetchall():
        rows.append({'id':'G'+str(g['id']),'service_date':g['service_date'],'branch_id':None,'exception_type':'CLOSED_ALL','reason':g['reason'],'created_at':g['created_at'],'branch_code':'ALL','branch_name':'全部分館','route_code':'全部'})
    c.close(); rows.sort(key=lambda x:(x['service_date'],str(x.get('route_code') or ''))); return rows

def date_span(start_date,end_date):
    try: a=date.fromisoformat(start_date); b=date.fromisoformat(end_date or start_date)
    except: raise HTTPException(400,'日期格式錯誤')
    if b<a: a,b=b,a
    if (b-a).days>62: raise HTTPException(400,'單次日期範圍最多63天')
    return [(a+timedelta(days=i)).isoformat() for i in range((b-a).days+1)]

@app.post('/api/schedule-exceptions/range')
async def save_schedule_range(req:Request):
    u=require_user(req,['SECRETARY']); p=await req.json(); et=(p.get('exception_type') or '').upper(); reason=(p.get('reason') or '').strip(); dates=date_span((p.get('start_date') or '').strip(),(p.get('end_date') or '').strip())
    if et not in ('CLOSED_ALL','STOP','ADD'): raise HTTPException(400,'類型錯誤')
    c=db()
    if et=='CLOSED_ALL':
        for d in dates:
            started=c.execute("SELECT 1 FROM deliveries WHERE service_date=? AND (document_original IS NOT NULL OR outbound_original IS NOT NULL OR branch_signed_at IS NOT NULL) LIMIT 1",(d,)).fetchone()
            if started: c.close(); raise HTTPException(409,f'{d} 已有配送開始，不能設定全館休館')
        for d in dates:
            c.execute("INSERT INTO global_closures(service_date,reason,created_by,created_at) VALUES(?,?,?,?) ON CONFLICT(service_date) DO UPDATE SET reason=excluded.reason,created_by=excluded.created_by,created_at=excluded.created_at",(d,reason,u['id'],now()))
            audit(c,'USER',u['id'],u['role'],'SET_GLOBAL_CLOSURE','GLOBAL_CLOSURE',d,after={'reason':reason})
            rebuild_service_date(c,d)
    else:
        bid=int(p.get('branch_id') or 0); b=c.execute('SELECT * FROM branches WHERE id=?',(bid,)).fetchone()
        if not b: c.close(); raise HTTPException(404,'找不到分館')
        if et=='STOP':
            for d in dates:
                x=c.execute('SELECT * FROM deliveries WHERE service_date=? AND branch_id=?',(d,bid)).fetchone()
                if x and (x['document_original'] is not None or x['outbound_original'] is not None or x['branch_signed_at'] is not None): c.close(); raise HTTPException(409,f'{d} 配送已開始，不能設定停送')
        for d in dates:
            c.execute("""INSERT INTO delivery_exceptions(service_date,branch_id,exception_type,reason,created_by,created_at) VALUES(?,?,?,?,?,?) ON CONFLICT(service_date,branch_id) DO UPDATE SET exception_type=excluded.exception_type,reason=excluded.reason,created_by=excluded.created_by,created_at=excluded.created_at""",(d,bid,et,reason,u['id'],now()))
            audit(c,'USER',u['id'],u['role'],'SET_DELIVERY_EXCEPTION','BRANCH',bid,after={'service_date':d,'exception_type':et,'reason':reason})
            rebuild_service_date(c,d)
    c.commit(); c.close(); _prefill_cache_clear(); await publish({'type':'schedule.updated','dates':dates}); return {'ok':True,'days':len(dates)}

@app.delete('/api/schedule-exceptions/{eid}')
async def delete_schedule_exception(eid:str,req:Request):
    u=require_user(req,['SECRETARY']); c=db()
    if eid.startswith('G'):
        try: gid=int(eid[1:])
        except: c.close(); raise HTTPException(400)
        x=c.execute('SELECT * FROM global_closures WHERE id=?',(gid,)).fetchone()
        if not x: c.close(); raise HTTPException(404)
        c.execute('DELETE FROM global_closures WHERE id=?',(gid,)); audit(c,'USER',u['id'],u['role'],'DELETE_GLOBAL_CLOSURE','GLOBAL_CLOSURE',x['service_date'],before=dict(x)); rebuild_service_date(c,x['service_date'])
    else:
        try: iid=int(eid)
        except: c.close(); raise HTTPException(400)
        x=c.execute('SELECT * FROM delivery_exceptions WHERE id=?',(iid,)).fetchone()
        if not x: c.close(); raise HTTPException(404)
        c.execute('DELETE FROM delivery_exceptions WHERE id=?',(iid,)); audit(c,'USER',u['id'],u['role'],'DELETE_DELIVERY_EXCEPTION','BRANCH',x['branch_id'],before=dict(x)); rebuild_service_date(c,x['service_date'])
    c.commit(); c.close(); _prefill_cache_clear(); await publish({'type':'schedule.updated'}); return {'ok':True}

@app.get('/api/secretary/sign-status')
def secretary_sign_status(req:Request):
    require_user(req,['ADMIN','SECRETARY']); c=db()
    rows=[dict(x) for x in c.execute("""SELECT dr.id,r.code,r.name,d.name driver_name,dr.status,dr.driver_signed_at,dr.driver_signature,dr.secretary_signature,dr.secretary_signed_at,(SELECT COUNT(*) FROM deliveries x WHERE x.daily_route_id=dr.id) total,(SELECT COUNT(*) FROM deliveries x WHERE x.daily_route_id=dr.id AND x.status='STOP_COMPLETED') completed,(SELECT COALESCE(SUM(x.document_final),0) FROM deliveries x WHERE x.daily_route_id=dr.id) document_total,(SELECT COALESCE(SUM(x.outbound_final),0) FROM deliveries x WHERE x.daily_route_id=dr.id) outbound_total,(SELECT COALESCE(SUM(x.inbound_final),0) FROM deliveries x WHERE x.daily_route_id=dr.id) inbound_total FROM daily_routes dr JOIN routes r ON r.id=dr.route_id LEFT JOIN drivers d ON d.id=dr.driver_id WHERE dr.service_date=? ORDER BY r.code""",(today(),)).fetchall()]
    report=c.execute('SELECT * FROM daily_reports WHERE service_date=?',(today(),)).fetchone(); c.close()
    active=[r for r in rows if r['total']>0]
    return {'routes':rows,'all_required_routes_signed':bool(active) and all(r['status']=='DRIVER_SIGNED' for r in active),'report':dict(report) if report else None}

@app.get('/api/routes/{rid}/signed-summary')
def route_signed_summary(rid:int,req:Request):
    require_user(req,['ADMIN','SECRETARY']); c=db()
    dr=c.execute("""SELECT dr.id,dr.service_date,dr.status,dr.driver_signed_at,dr.driver_signature,dr.secretary_signature,dr.secretary_signed_at,dr.secretary_signature,dr.secretary_signed_at,r.code,r.name,d.name driver_name FROM daily_routes dr JOIN routes r ON r.id=dr.route_id LEFT JOIN drivers d ON d.id=dr.driver_id WHERE dr.id=?""",(rid,)).fetchone()
    if not dr: c.close(); raise HTTPException(404)
    rows=[dict(x) for x in c.execute("""SELECT b.code branch_code,b.name branch_name,x.document_final,x.outbound_final,x.inbound_final,x.note_final,x.signer_name,x.correction_signer_name,x.branch_signed_at,x.corrected_at,x.branch_signature,x.correction_signature,x.status FROM deliveries x JOIN branches b ON b.id=x.branch_id WHERE x.daily_route_id=? ORDER BY b.stop_order""",(rid,)).fetchall()]
    totals={'document':sum((x['document_final'] or 0) for x in rows),'outbound':sum((x['outbound_final'] or 0) for x in rows),'inbound':sum((x['inbound_final'] or 0) for x in rows)}
    c.close(); return {'route':dict(dr),'stops':rows,'totals':totals}


@app.post('/api/routes/{rid}/secretary-sign')
async def secretary_route_sign(rid:int,req:Request):
    u=require_user(req,['SECRETARY']); p=await req.json(); sig=p.get('signature','')
    if not sig: raise HTTPException(400,'秘書簽名必填')
    c=db(); dr=c.execute("SELECT * FROM daily_routes WHERE id=?",(rid,)).fetchone()
    if not dr: c.close(); raise HTTPException(404)
    if dr['status']!='DRIVER_SIGNED': c.close(); raise HTTPException(409,'需先完成司機路線簽名')
    c.execute('UPDATE daily_routes SET secretary_signature=?,secretary_signed_at=? WHERE id=?',(sig,now(),rid))
    audit(c,'USER',u['id'],u['role'],'SECRETARY_ROUTE_SIGN','DAILY_ROUTE',rid,after={'signed_at':now()})
    c.commit(); c.close(); await publish({'type':'route.secretary_signed','id':rid}); return {'ok':True}

@app.post('/api/secretary/final-sign')
async def final_sign(req:Request):
    u=require_user(req,['SECRETARY']);p=await req.json(); sig=p.get('signature','')
    if not sig: raise HTTPException(400,'秘書簽名必填')
    c=db(); active=c.execute("""SELECT dr.id,dr.status FROM daily_routes dr WHERE dr.service_date=? AND EXISTS(SELECT 1 FROM deliveries x WHERE x.daily_route_id=dr.id)""",(today(),)).fetchall()
    if not active: c.close(); raise HTTPException(409,'今日沒有配送路線')
    if any(x['status']!='DRIVER_SIGNED' for x in active): c.close(); raise HTTPException(409,'尚有配送路線未完成司機簽名')
    if is_report_locked(c): c.close(); raise HTTPException(409,'今日日報已完成簽名並鎖定')
    c.execute("INSERT INTO daily_reports(service_date,secretary_signature,secretary_signed_at,status,locked_at) VALUES(?,?,?,?,?) ON CONFLICT(service_date) DO UPDATE SET secretary_signature=excluded.secretary_signature,secretary_signed_at=excluded.secretary_signed_at,status='LOCKED',locked_at=excluded.locked_at",(today(),sig,now(),'LOCKED',now()))
    audit(c,'USER',u['id'],u['role'],'SECRETARY_FINAL_SIGN','DAY',today(),after={'locked_at':now()}); c.commit();c.close();await publish({'type':'report.locked'});return {'ok':True}

@app.get('/api/reports/today.csv')
def report_csv(req:Request):
    require_user(req,['ADMIN','SECRETARY']);c=db();rows=c.execute('''SELECT r.code 路線,d.name 司機,b.name 分館,x.document_final 公文,x.outbound_final 圖書送出,x.inbound_final 圖書收回,x.note_final 備註,x.signer_name 簽收人,x.branch_signed_at 簽收時間,x.status 狀態 FROM deliveries x JOIN branches b ON b.id=x.branch_id JOIN daily_routes dr ON dr.id=x.daily_route_id JOIN routes r ON r.id=dr.route_id JOIN drivers d ON d.id=dr.driver_id WHERE x.service_date=? ORDER BY r.code,b.stop_order''',(today(),)).fetchall();c.close();bio=io.StringIO();w=csv.writer(bio);headers=rows[0].keys() if rows else [];w.writerow(headers);[w.writerow(list(r)) for r in rows];data='\ufeff'+bio.getvalue();return Response(data,media_type='text/csv',headers={'Content-Disposition':f'attachment; filename="{today()}_library_logistics.csv"'})
@app.get('/api/audit')
def audits(req:Request,limit:int=100):
    require_user(req,['ADMIN','SECRETARY']);c=db();rows=[dict(x) for x in c.execute('SELECT * FROM audit_logs ORDER BY id DESC LIMIT ?',(min(limit,500),)).fetchall()];c.close();return rows


# ===== Expanded Admin / Secretary management =====
@app.post('/api/branches')
async def create_branch(req:Request):
    u=require_user(req,['ADMIN','SECRETARY']); p=await req.json(); c=db(); code=(p.get('code') or '').strip(); name=(p.get('name') or '').strip()
    if not code or not name: c.close(); raise HTTPException(400,'分館代碼與名稱必填')
    if c.execute('SELECT 1 FROM branches WHERE code=?',(code,)).fetchone(): c.close(); raise HTTPException(409,'分館代碼已存在')
    route_id=int(p.get('route_id') or 1); stop_order=int(p.get('stop_order') or 1); raw=secrets.token_urlsafe(24)
    c.execute('''INSERT INTO branches(code,name,route_id,stop_order,active,pin_hash,access_token_hash,qr_created_at,address,phone,contact_name,contact_info,delivery_weekdays,delivery_frequency) VALUES(?,?,?,?,1,?,?,?,?,?,?,?,?,?)''',(code,name,route_id,stop_order,hash_secret(p.get('pin') or '1234'),thash(raw),now(),p.get('address',''),p.get('phone',''),p.get('contact_name',''),p.get('contact_info',''),p.get('delivery_weekdays','1,2,3,4,5'),p.get('delivery_frequency','每週固定')))
    bid=c.execute('SELECT last_insert_rowid()').fetchone()[0]; audit(c,'USER',u['id'],u['role'],'CREATE_BRANCH','BRANCH',bid,after={'token':raw,'code':code,'name':name}); c.commit(); c.close(); return {'ok':True,'id':bid}

@app.get('/api/branches/{bid}/detail')
def branch_detail(req:Request,bid:int):
    require_user(req,['ADMIN','SECRETARY']); c=db(); b=c.execute('SELECT b.*,r.code route_code FROM branches b LEFT JOIN routes r ON r.id=b.route_id WHERE b.id=?',(bid,)).fetchone(); c.close()
    if not b: raise HTTPException(404)
    out=dict(b); out['pin_status']='已設定' if out.get('pin_hash') else '未設定'; out.pop('pin_hash',None); out.pop('access_token_hash',None); return out

@app.patch('/api/branches/{bid}')
async def update_branch(req:Request,bid:int):
    u=require_user(req,['ADMIN','SECRETARY']); p=await req.json(); c=db(); old=c.execute('SELECT * FROM branches WHERE id=?',(bid,)).fetchone()
    if not old: c.close(); raise HTTPException(404)
    fields=['code','name','route_id','stop_order','address','phone','contact_name','contact_info','delivery_weekdays','delivery_frequency']; vals={k:p[k] for k in fields if k in p}
    if vals: c.execute('UPDATE branches SET '+','.join(k+'=?' for k in vals)+' WHERE id=?',(*vals.values(),bid))
    audit(c,'USER',u['id'],u['role'],'UPDATE_BRANCH','BRANCH',bid,after=vals); c.commit(); c.close(); return {'ok':True}

@app.post('/api/branches/{bid}/deactivate')
async def deactivate_branch(req:Request,bid:int):
    u=require_user(req,['ADMIN','SECRETARY']); c=db(); c.execute('UPDATE branches SET active=0 WHERE id=?',(bid,)); audit(c,'USER',u['id'],u['role'],'DEACTIVATE_BRANCH','BRANCH',bid); c.commit(); c.close(); return {'ok':True}
@app.post('/api/branches/{bid}/activate')
async def activate_branch(req:Request,bid:int):
    u=require_user(req,['ADMIN','SECRETARY']); c=db(); c.execute('UPDATE branches SET active=1 WHERE id=?',(bid,)); audit(c,'USER',u['id'],u['role'],'ACTIVATE_BRANCH','BRANCH',bid); c.commit(); c.close(); return {'ok':True}

@app.delete('/api/branches/{bid}')
async def delete_branch(req:Request,bid:int):
    u=require_user(req,['ADMIN']); c=db(); b=c.execute('SELECT * FROM branches WHERE id=?',(bid,)).fetchone()
    if not b: c.close(); raise HTTPException(404,'找不到分館')
    used=c.execute('SELECT 1 FROM deliveries WHERE branch_id=? LIMIT 1',(bid,)).fetchone()
    if used: c.close(); raise HTTPException(409,'此分館已有配送歷史，為保留簽收紀錄不可永久刪除；請改用「停用」')
    c.execute('DELETE FROM branches WHERE id=?',(bid,)); audit(c,'USER',u['id'],u['role'],'DELETE_BRANCH','BRANCH',bid,after={'code':b['code'],'name':b['name']}); c.commit(); c.close(); return {'ok':True}

@app.get('/api/routes')
def list_routes(req:Request):
    require_user(req,['ADMIN','SECRETARY']); c=db(); rows=[dict(x) for x in c.execute('SELECT * FROM routes ORDER BY code').fetchall()]; c.close(); return rows
@app.post('/api/routes')
async def create_route(req:Request):
    u=require_user(req,['ADMIN']); p=await req.json(); c=db(); cur=c.execute('INSERT INTO routes(code,name,active) VALUES(?,?,1)',((p.get('code') or '').strip(),(p.get('name') or '').strip())); rid=cur.lastrowid; audit(c,'USER',u['id'],u['role'],'CREATE_ROUTE','ROUTE',rid); c.commit(); c.close(); return {'ok':True}
@app.patch('/api/routes/{rid}')
async def update_route(req:Request,rid:int):
    u=require_user(req,['ADMIN']); p=await req.json(); c=db(); vals={k:p[k] for k in ['code','name','active'] if k in p}
    if vals:c.execute('UPDATE routes SET '+','.join(k+'=?' for k in vals)+' WHERE id=?',(*vals.values(),rid))
    audit(c,'USER',u['id'],u['role'],'UPDATE_ROUTE','ROUTE',rid,after=vals); c.commit(); c.close(); return {'ok':True}

@app.post('/api/drivers')
async def create_driver(req:Request):
    u=require_user(req,['ADMIN']); p=await req.json(); name=(p.get('name') or '').strip()
    if not name: raise HTTPException(400,'司機姓名必填')
    c=db(); cur=c.execute('INSERT INTO drivers(name,active) VALUES(?,1)',(name,)); did=cur.lastrowid; audit(c,'USER',u['id'],u['role'],'CREATE_DRIVER','DRIVER',did); c.commit(); c.close(); return {'ok':True,'id':did}
@app.patch('/api/drivers/{did}')
async def update_driver(req:Request,did:int):
    u=require_user(req,['ADMIN']); p=await req.json(); c=db(); vals={k:p[k] for k in ['name','active'] if k in p}
    if vals:c.execute('UPDATE drivers SET '+','.join(k+'=?' for k in vals)+' WHERE id=?',(*vals.values(),did))
    audit(c,'USER',u['id'],u['role'],'UPDATE_DRIVER','DRIVER',did,after=vals); c.commit(); c.close(); return {'ok':True}

@app.delete('/api/drivers/{did}')
async def delete_driver(req:Request,did:int):
    u=require_user(req,['ADMIN']); c=db(); d=c.execute('SELECT * FROM drivers WHERE id=?',(did,)).fetchone()
    if not d: c.close(); raise HTTPException(404,'找不到司機')
    used=c.execute('SELECT 1 FROM daily_routes WHERE driver_id=? LIMIT 1',(did,)).fetchone()
    devices=c.execute('SELECT 1 FROM driver_devices WHERE driver_id=? LIMIT 1',(did,)).fetchone()
    if used or devices: c.close(); raise HTTPException(409,'此司機已有配送或裝置歷史，為保留紀錄不可永久刪除；請改用「停用」')
    c.execute('DELETE FROM drivers WHERE id=?',(did,)); audit(c,'USER',u['id'],u['role'],'DELETE_DRIVER','DRIVER',did,after={'name':d['name']}); c.commit(); c.close(); return {'ok':True}
@app.post('/api/drivers/{did}/pin')
async def reset_driver_pin(req:Request,did:int):
    u=require_user(req,['ADMIN']); p=await req.json(); pin=str(p.get('pin',''))
    if len(pin)!=4 or not pin.isdigit(): raise HTTPException(400,'PIN需為4位數字')
    c=db(); c.execute('UPDATE drivers SET pin_hash=? WHERE id=?',(hash_secret(pin),did)); audit(c,'USER',u['id'],u['role'],'RESET_DRIVER_PIN','DRIVER',did); c.commit(); c.close(); return {'ok':True}
@app.get('/api/drivers/{did}/devices')
def driver_devices(req:Request,did:int):
    require_user(req,['ADMIN']); c=db(); rows=[dict(x) for x in c.execute('SELECT id,driver_id,activated_at,revoked_at,last_seen_at FROM driver_devices WHERE driver_id=? ORDER BY id DESC',(did,)).fetchall()]; c.close(); return rows
@app.post('/api/driver-devices/{device_id}/revoke')
async def revoke_device(req:Request,device_id:int):
    u=require_user(req,['ADMIN']); c=db(); c.execute('UPDATE driver_devices SET revoked_at=? WHERE id=?',(now(),device_id)); audit(c,'USER',u['id'],u['role'],'REVOKE_DRIVER_DEVICE','DRIVER_DEVICE',device_id); c.commit(); c.close(); return {'ok':True}

@app.get('/api/daily-routes')
def daily_route_list(req:Request, service_date:str|None=None):
    require_user(req,['ADMIN','SECRETARY']); d=service_date or today(); c=db()
    rows=[dict(x) for x in c.execute('''SELECT dr.id,dr.service_date,dr.status,
        r.id route_id,r.code,r.name,dv.id driver_id,dv.name driver_name,
        (SELECT COUNT(*) FROM deliveries x WHERE x.daily_route_id=dr.id) total,
        (SELECT COUNT(*) FROM deliveries x WHERE x.daily_route_id=dr.id AND x.status='STOP_COMPLETED') completed
        FROM daily_routes dr
        JOIN routes r ON r.id=dr.route_id
        LEFT JOIN drivers dv ON dv.id=dr.driver_id
        WHERE dr.service_date=?
        ORDER BY CAST(r.code AS INTEGER),r.code''',(d,)).fetchall()]
    c.close(); return rows
@app.patch('/api/daily-routes/{drid}/driver')
async def assign_driver(req:Request,drid:int):
    u=require_user(req,['ADMIN']); p=await req.json(); did=int(p.get('driver_id')); c=db(); c.execute('UPDATE daily_routes SET driver_id=? WHERE id=?',(did,drid)); audit(c,'USER',u['id'],u['role'],'ASSIGN_DAILY_DRIVER','DAILY_ROUTE',drid,after={'driver_id':did}); c.commit(); c.close(); await publish({'type':'route.updated','id':drid}); return {'ok':True}

@app.get('/api/deliveries/all')
def all_deliveries(req:Request, service_date:str|None=None):
    require_user(req,['ADMIN','SECRETARY']); d=service_date or today(); c=db(); rows=[dict(x) for x in c.execute('''SELECT x.*,b.code branch_code,b.name branch_name,r.code route_code,d.name driver_name,CASE WHEN co.id IS NULL THEN 0 ELSE 1 END has_correction,co.driver_note correction_driver_note,co.status correction_status FROM deliveries x JOIN branches b ON b.id=x.branch_id JOIN daily_routes dr ON dr.id=x.daily_route_id JOIN routes r ON r.id=dr.route_id JOIN drivers d ON d.id=dr.driver_id LEFT JOIN corrections co ON co.delivery_id=x.id WHERE x.service_date=? ORDER BY r.code,b.stop_order''',(d,)).fetchall()]; c.close(); return rows
@app.get('/api/corrections')
def correction_list(req:Request):
    require_user(req,['ADMIN','SECRETARY']); c=db(); rows=[dict(x) for x in c.execute('''SELECT co.*,b.name branch_name,r.code route_code,d.name driver_name,x.document_final,x.outbound_final,x.inbound_final,x.note_final,x.correction_reason,x.correction_signer_name,x.corrected_at FROM corrections co JOIN deliveries x ON x.id=co.delivery_id JOIN branches b ON b.id=x.branch_id JOIN daily_routes dr ON dr.id=x.daily_route_id JOIN routes r ON r.id=dr.route_id JOIN drivers d ON d.id=co.requested_by_driver_id ORDER BY co.id DESC''').fetchall()]; c.close(); return rows

def report_rows(c,start_date,end_date):
    rows=[dict(x) for x in c.execute("""SELECT x.service_date 日期,r.code 路線,COALESCE(ad.name,dv.name,'') 實際配送司機,b.code 分館代碼,b.name 分館,x.document_final 公文,x.outbound_final 圖書送出,x.inbound_final 圖書收回,COALESCE(x.note_final,'') 備註,COALESCE(x.correction_signer_name,x.signer_name,'') 簽收人,COALESCE(x.corrected_at,x.branch_signed_at,'') 簽收時間,CASE WHEN EXISTS(SELECT 1 FROM corrections co2 WHERE co2.delivery_id=x.id) THEN '是' ELSE '否' END 是否更正,COALESCE(x.correction_reason,'') 更正原因 FROM deliveries x JOIN branches b ON b.id=x.branch_id JOIN daily_routes dr ON dr.id=x.daily_route_id JOIN routes r ON r.id=dr.route_id LEFT JOIN delivery_driver_assignments dda ON dda.delivery_id=x.id LEFT JOIN drivers ad ON ad.id=dda.driver_id LEFT JOIN drivers dv ON dv.id=dr.driver_id WHERE x.service_date BETWEEN ? AND ? ORDER BY x.service_date,r.code,b.stop_order""",(start_date,end_date)).fetchall()]
    for r in rows: r['簽收時間']=report_time(r.get('簽收時間'))
    return rows

def decode_data_url(v):
    if not v or ',' not in v: return None
    try: return base64.b64decode(v.split(',',1)[1])
    except: return None

def official_daily_payload(c,d):
    rep=c.execute('SELECT * FROM daily_reports WHERE service_date=?',(d,)).fetchone()
    routes=[dict(x) for x in c.execute("""SELECT dr.id,r.code,dv.name driver_name,dr.driver_signature,dr.driver_signed_at FROM daily_routes dr JOIN routes r ON r.id=dr.route_id LEFT JOIN drivers dv ON dv.id=dr.driver_id WHERE dr.service_date=? AND EXISTS(SELECT 1 FROM deliveries x WHERE x.daily_route_id=dr.id) ORDER BY r.code""",(d,)).fetchall()]
    for r in routes:
        r['stops']=[dict(x) for x in c.execute("""SELECT b.code branch_code,b.name branch_name,x.document_final,x.outbound_final,x.inbound_final,x.note_final,x.signer_name,x.correction_signer_name,x.branch_signed_at,x.corrected_at,x.branch_signature,x.correction_signature,x.status FROM deliveries x JOIN branches b ON b.id=x.branch_id WHERE x.daily_route_id=? ORDER BY b.stop_order""",(r['id'],)).fetchall()]
        r['totals']={'document':sum((x['document_final'] or 0) for x in r['stops']),'outbound':sum((x['outbound_final'] or 0) for x in r['stops']),'inbound':sum((x['inbound_final'] or 0) for x in r['stops'])}
    grand={'document':sum(r['totals']['document'] for r in routes),'outbound':sum(r['totals']['outbound'] for r in routes),'inbound':sum(r['totals']['inbound'] for r in routes)}
    return {'report':dict(rep) if rep else None,'routes':routes,'grand':grand}

def assert_official_locked(c,d):
    r=c.execute("SELECT status FROM daily_reports WHERE service_date=?",(d,)).fetchone()
    if not r or r['status']!='LOCKED': raise HTTPException(409,'正式日報需完成秘書最終簽名並鎖定後才能下載')

def xlsx_bytes(rows,title):
    wb=Workbook(); ws=wb.active; ws.title='配送報表'; ws.append([title])
    if rows:
        headers=list(rows[0].keys()); ws.append(headers)
        for r in rows: ws.append([r[h] for h in headers])
        ws.freeze_panes='A3'
        for col in ws.columns:
            ws.column_dimensions[col[0].column_letter].width=min(max(12,max(len(str(c.value or '')) for c in col)+2),30)
    bio=io.BytesIO(); wb.save(bio); return bio.getvalue()

def official_xlsx(payload,d):
    wb=Workbook(); ws=wb.active; ws.title='今日路線總表'
    ws.append([f'{d} 圖書物流配送暨電子簽收日報']); ws.append(['路線','司機','分館','公文','圖書送出','圖書收回','簽收人','簽收時間','備註'])
    row=3
    for r in payload['routes']:
        for x in r['stops']:
            ws.append([r['code'],r['driver_name'],x['branch_name'],x['document_final'] or 0,x['outbound_final'] or 0,x['inbound_final'] or 0,x['correction_signer_name'] or x['signer_name'] or '',x['corrected_at'] or x['branch_signed_at'] or '',x['note_final'] or '']); row+=1
        ws.append([f"{r['code']}線合計",r['driver_name'],'',r['totals']['document'],r['totals']['outbound'],r['totals']['inbound'],'','','']); row+=1
    ws.append(['全部總計','','',payload['grand']['document'],payload['grand']['outbound'],payload['grand']['inbound'],'','','']); row+=2
    ws.append(['司機簽名']); row+=1
    for r in payload['routes']:
        ws.append([f"{r['code']}線 {r['driver_name'] or ''}",r['driver_signed_at'] or ''])
        data=decode_data_url(r.get('driver_signature'))
        if data:
            try:
                img=XLImage(io.BytesIO(data)); img.width=180; img.height=55; ws.add_image(img,f'C{row}')
            except: pass
        row+=4
    rep=payload['report'] or {}; ws.append(['總館秘書最終簽名',rep.get('secretary_signed_at') or ''])
    data=decode_data_url(rep.get('secretary_signature'))
    if data:
        try:
            img=XLImage(io.BytesIO(data)); img.width=200; img.height=65; ws.add_image(img,f'C{row}')
        except: pass
    for col,w in {'A':16,'B':16,'C':22,'D':12,'E':14,'F':14,'G':16,'H':22,'I':24}.items(): ws.column_dimensions[col].width=w
    bio=io.BytesIO(); wb.save(bio); return bio.getvalue()


def register_pdf_zh_font():
    # Prefer an embedded Traditional-Chinese TrueType font so downloaded PDFs render
    # correctly in Safari/Preview/Chrome instead of relying on viewer CID substitution.
    candidates=[
        '/usr/share/fonts/truetype/arphic/uming.ttc',
        '/usr/share/fonts/truetype/arphic-bkai00mp/bkai00mp.ttf',
    ]
    for fp in candidates:
        if os.path.exists(fp):
            try:
                pdfmetrics.registerFont(TTFont('ReportZH',fp,subfontIndex=0))
                return 'ReportZH'
            except Exception:
                pass
    pdfmetrics.registerFont(UnicodeCIDFont('MSung-Light'))
    return 'MSung-Light'

def pdf_bytes(rows,title):
    bio=io.BytesIO(); zh=register_pdf_zh_font(); cv=canvas.Canvas(bio,pagesize=(842,595)); cv.setFont(zh,16); cv.drawString(30,565,title); y=540; cv.setFont(zh,7.3)
    for r in rows:
        line=f"{r['日期']}  {r['路線']}線  {r['分館代碼']} {r['分館']}  公文:{r['公文'] or 0}  送出:{r['圖書送出'] or 0}  收回:{r['圖書收回'] or 0}  簽收:{r['簽收人'] or '—'}  更正:{r['是否更正']}"; cv.drawString(25,y,line[:150]); y-=12
        extra=f"備註:{r['備註'] or '—'}  簽收時間:{r['簽收時間'] or '—'}  更正原因:{r['更正原因'] or '—'}"; cv.drawString(40,y,extra[:160]); y-=14
        if y<30: cv.showPage(); cv.setFont(zh,7.3); y=565
    cv.save(); return bio.getvalue()

def official_pdf(payload,d):
    bio=io.BytesIO(); zh=register_pdf_zh_font(); cv=canvas.Canvas(bio,pagesize=(842,595))
    cv.setFont(zh,16); cv.drawString(30,565,f'{d} 圖書物流配送暨電子簽收日報'); y=538
    for r in payload['routes']:
        if y<110: cv.showPage(); y=565
        cv.setFont(zh,12); cv.drawString(30,y,f"{r['code']}線｜司機 {r['driver_name'] or '—'}"); y-=18; cv.setFont(zh,8)
        for x in r['stops']:
            cv.drawString(42,y,f"{x['branch_name']}  公文 {x['document_final'] or 0}｜送出 {x['outbound_final'] or 0}｜收回 {x['inbound_final'] or 0}｜簽收 {x['correction_signer_name'] or x['signer_name'] or '—'}"); y-=13
            if y<90: cv.showPage(); cv.setFont(zh,8); y=565
        cv.setFont(zh,9); cv.drawString(42,y,f"路線合計：公文 {r['totals']['document']}｜送出 {r['totals']['outbound']}｜收回 {r['totals']['inbound']}"); y-=16
        data=decode_data_url(r.get('driver_signature'))
        if data:
            try:
                cv.drawImage(ImageReader(io.BytesIO(data)),42,y-45,width=150,height=42,mask='auto'); cv.drawString(200,y-22,f"司機簽名 {r['driver_signed_at'] or ''}"); y-=55
            except: pass
    if y<110: cv.showPage(); y=565
    cv.setFont(zh,11); cv.drawString(30,y,f"全部總計：公文 {payload['grand']['document']}｜送出 {payload['grand']['outbound']}｜收回 {payload['grand']['inbound']}"); y-=22
    rep=payload['report'] or {}; data=decode_data_url(rep.get('secretary_signature'))
    if data:
        try: cv.drawImage(ImageReader(io.BytesIO(data)),30,y-55,width=180,height=50,mask='auto')
        except: pass
    cv.setFont(zh,9); cv.drawString(225,y-28,f"總館秘書最終簽名 {rep.get('secretary_signed_at') or ''}")
    cv.save(); return bio.getvalue()

@app.get('/api/reports/daily.xlsx')
def daily_xlsx(req:Request,service_date:str|None=None):
    require_user(req,['ADMIN','SECRETARY']); d=service_date or today(); c=db(); rows=report_rows(c,d,d); c.close(); return Response(xlsx_bytes(rows,f'{d} 圖書物流配送日報'),media_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',headers={'Content-Disposition':f'attachment; filename="{d}_daily.xlsx"'})
@app.get('/api/reports/daily.pdf')
def daily_pdf(req:Request,service_date:str|None=None):
    require_user(req,['ADMIN','SECRETARY']); d=service_date or today(); c=db(); rows=report_rows(c,d,d); c.close(); return Response(pdf_bytes(rows,f'{d} 圖書物流配送日報'),media_type='application/pdf',headers={'Content-Disposition':f'attachment; filename="{d}_daily.pdf"'})
def month_range(m):
    start=m+'-01'; y=int(m[:4]); mo=int(m[5:7]); nxt=f'{y+1}-01-01' if mo==12 else f'{y}-{mo+1:02d}-01'; end=(datetime.fromisoformat(nxt)-timedelta(days=1)).date().isoformat(); return start,end
@app.get('/api/reports/monthly.xlsx')
def monthly_xlsx(req:Request,month:str|None=None):
    require_user(req,['ADMIN','SECRETARY']); m=month or today()[:7]; a,b=month_range(m); c=db(); rows=report_rows(c,a,b); c.close(); return Response(xlsx_bytes(rows,f'{m} 圖書物流配送月報'),media_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',headers={'Content-Disposition':f'attachment; filename="{m}_monthly.xlsx"'})
@app.get('/api/reports/monthly.pdf')
def monthly_pdf(req:Request,month:str|None=None):
    require_user(req,['ADMIN','SECRETARY']); m=month or today()[:7]; a,b=month_range(m); c=db(); rows=report_rows(c,a,b); c.close(); return Response(pdf_bytes(rows,f'{m} 圖書物流配送月報'),media_type='application/pdf',headers={'Content-Disposition':f'attachment; filename="{m}_monthly.pdf"'})

@app.get('/api/email/settings')
def email_settings(req:Request):
    require_user(req,['ADMIN']); cfg=smtp_config(); diag=smtp_diagnostics(); return {'host':cfg['host'],'port':cfg['port'],'user':cfg['user'],'from':cfg['from'],'tls':cfg['tls'],'password_configured':bool(cfg['password']),'checks':diag['checks'],'source':diag['source'],'missing':diag['missing'],'configured':not bool(diag['missing'])}

@app.post('/api/email/test')
async def email_test(req:Request):
    require_user(req,['ADMIN']); p=await req.json(); to=(p.get('email') or '').strip()
    if '@' not in to: raise HTTPException(400,'請輸入測試收件信箱')
    try: smtp_send('圖書物流系統 SMTP 測試','Google Workspace SMTP 已成功連線。\n\n寄件者：'+smtp_config()['from'],[to])
    except Exception as e: raise HTTPException(500,str(e))
    return {'ok':True,'message':'測試郵件已寄出'}

@app.get('/api/email/recipients')
def recipients(req:Request):
    require_user(req,['ADMIN','SECRETARY']); c=db(); rows=[dict(x) for x in c.execute('SELECT * FROM email_recipients ORDER BY id').fetchall()]; c.close(); return rows
@app.post('/api/email/recipients')
async def add_recipient(req:Request):
    u=require_user(req,['ADMIN']); p=await req.json(); email=(p.get('email') or '').strip(); typ=(p.get('recipient_type') or 'TO').upper()
    if '@' not in email: raise HTTPException(400,'Email格式錯誤')
    if typ not in ('TO','CC'): raise HTTPException(400,'收件類型錯誤')
    c=db(); cur=c.execute('INSERT INTO email_recipients(email,recipient_type,active,created_at) VALUES(?,?,1,?)',(email,typ,now())); rid=cur.lastrowid; audit(c,'USER',u['id'],u['role'],'ADD_EMAIL_RECIPIENT','EMAIL_RECIPIENT',rid); c.commit(); c.close(); return {'ok':True}
@app.post('/api/email/recipients/{rid}/toggle')
async def toggle_recipient(req:Request,rid:int):
    u=require_user(req,['ADMIN']); c=db(); c.execute('UPDATE email_recipients SET active=CASE active WHEN 1 THEN 0 ELSE 1 END WHERE id=?',(rid,)); audit(c,'USER',u['id'],u['role'],'TOGGLE_EMAIL_RECIPIENT','EMAIL_RECIPIENT',rid); c.commit(); c.close(); return {'ok':True}
@app.delete('/api/email/recipients/{rid}')
def delete_recipient(req:Request,rid:int):
    u=require_user(req,['ADMIN']); c=db(); c.execute('DELETE FROM email_recipients WHERE id=?',(rid,)); audit(c,'USER',u['id'],u['role'],'DELETE_EMAIL_RECIPIENT','EMAIL_RECIPIENT',rid); c.commit(); c.close(); return {'ok':True}
@app.get('/api/email/logs')
def email_logs(req:Request):
    require_user(req,['ADMIN','SECRETARY']); c=db(); rows=[dict(x) for x in c.execute('SELECT * FROM email_logs ORDER BY id DESC LIMIT 100').fetchall()]; c.close(); return rows

def build_report_email(report_type,period):
    c=db()
    if report_type=='MONTHLY':
        a,b=month_range(period[:7]); rows=report_rows(c,a,b); c.close(); title=f'{period[:7]} 圖書物流配送月報'; x=xlsx_bytes(rows,title); p=pdf_bytes(rows,title); base=period[:7]+'_monthly'
    else:
        d=period[:10]; rows=report_rows(c,d,d); c.close(); title=f'{d} 圖書物流配送日報'; x=xlsx_bytes(rows,title); p=pdf_bytes(rows,title); base=d+'_daily'
    return title,[(base+'.xlsx',x,'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'),(base+'.pdf',p,'application/pdf')]

@app.post('/api/email/send-report')
async def send_report(req:Request):
    u=require_user(req,['ADMIN','SECRETARY']); p=await req.json(); period=p.get('period') or today(); typ=(p.get('report_type') or 'DAILY').upper(); c=db(); rec=[dict(x) for x in c.execute('SELECT email,recipient_type FROM email_recipients WHERE active=1 ORDER BY id').fetchall()]
    if not rec: c.close(); raise HTTPException(409,'尚未設定Email收件人')
    to=[x['email'] for x in rec if x['recipient_type']=='TO']; cc=[x['email'] for x in rec if x['recipient_type']=='CC'];
    if not to and cc: to=[cc.pop(0)]
    cur=c.execute('INSERT INTO email_logs(report_type,period,recipients,status,created_at) VALUES(?,?,?,?,?)',(typ,period,','.join([x['email'] for x in rec]),'SENDING',now())); lid=cur.lastrowid; c.commit(); c.close()
    try:
        title,atts=build_report_email(typ,period); smtp_send(title,f'附件為 {period} 圖書物流報表。\n\n寄件者：{smtp_config()["from"]}',to,cc,atts)
        c=db(); c.execute("UPDATE email_logs SET status='SENT',sent_at=?,error_message=NULL WHERE id=?",(now(),lid)); audit(c,'USER',u['id'],u['role'],'SEND_REPORT_EMAIL','EMAIL_LOG',lid,after={'recipients':[x['email'] for x in rec]}); c.commit(); c.close(); return {'ok':True,'status':'SENT','recipients':[x['email'] for x in rec]}
    except Exception as e:
        c=db(); c.execute("UPDATE email_logs SET status='FAILED',error_message=? WHERE id=?",(str(e),lid)); c.commit(); c.close(); raise HTTPException(500,str(e))
@app.post('/api/email/logs/{lid}/resend')
async def resend_email(req:Request,lid:int):
    u=require_user(req,['ADMIN','SECRETARY']); c=db(); old=c.execute('SELECT * FROM email_logs WHERE id=?',(lid,)).fetchone(); c.close()
    if not old: raise HTTPException(404)
    rec=[x.strip() for x in (old['recipients'] or '').split(',') if x.strip()];
    if not rec: raise HTTPException(409,'原寄送紀錄沒有收件人')
    try:
        title,atts=build_report_email(old['report_type'],old['period']); smtp_send(title,f'重新寄送：附件為 {old["period"]} 圖書物流報表。',[rec[0]],rec[1:],atts)
        c=db(); c.execute("UPDATE email_logs SET status='SENT',sent_at=?,error_message=NULL WHERE id=?",(now(),lid)); audit(c,'USER',u['id'],u['role'],'RESEND_REPORT_EMAIL','EMAIL_LOG',lid); c.commit(); c.close(); return {'ok':True,'status':'SENT'}
    except Exception as e:
        c=db(); c.execute("UPDATE email_logs SET status='FAILED',error_message=? WHERE id=?",(str(e),lid)); c.commit(); c.close(); raise HTTPException(500,str(e))

@app.get('/api/events')
async def events(req:Request):
    require_user(req,['ADMIN','SECRETARY']);q=asyncio.Queue();subscribers.add(q)
    async def gen():
        try:
            while True:
                try:e=await asyncio.wait_for(q.get(),20);yield 'event: update\ndata: '+json.dumps(e,ensure_ascii=False)+'\n\n'
                except asyncio.TimeoutError:yield 'event: ping\ndata: {}\n\n'
        finally:subscribers.discard(q)
    return StreamingResponse(gen(),media_type='text/event-stream')


# V17 臨時交接 / 分段簽名 / SMTP
@app.get('/api/route-handoffs')
def route_handoffs(req:Request, service_date:str|None=None):
    require_user(req,['ADMIN','SECRETARY']); d=service_date or today(); c=db(); rows=[dict(x) for x in c.execute("SELECT h.*,r.code route_code,fd.name from_driver,td.name to_driver,b.name start_branch FROM route_handoffs h JOIN daily_routes dr ON dr.id=h.daily_route_id JOIN routes r ON r.id=dr.route_id JOIN drivers fd ON fd.id=h.from_driver_id JOIN drivers td ON td.id=h.to_driver_id JOIN deliveries x ON x.id=h.start_delivery_id JOIN branches b ON b.id=x.branch_id WHERE h.service_date=? ORDER BY h.id DESC",(d,)).fetchall()];c.close();return rows

@app.post('/api/daily-routes/{rid}/handoff')
async def route_handoff(rid:int,req:Request):
    u=require_user(req,['ADMIN']); p=await req.json(); to_driver=int(p.get('to_driver_id')); start_id=int(p.get('start_delivery_id')); reason=(p.get('reason') or '').strip(); note=(p.get('note') or '').strip()
    if not reason: raise HTTPException(400,'交接原因必填')
    c=db(); dr=c.execute('SELECT * FROM daily_routes WHERE id=?',(rid,)).fetchone(); start=c.execute('SELECT x.*,b.stop_order FROM deliveries x JOIN branches b ON b.id=x.branch_id WHERE x.id=? AND x.daily_route_id=?',(start_id,rid)).fetchone()
    if not dr or not start: c.close(); raise HTTPException(404,'找不到路線或起始分館')
    if start['status']=='STOP_COMPLETED': c.close(); raise HTTPException(409,'只能從尚未完成的分館開始交接')
    htime=now(); cur=c.execute('INSERT INTO route_handoffs(service_date,daily_route_id,from_driver_id,to_driver_id,start_delivery_id,reason,note,created_by,created_at) VALUES(?,?,?,?,?,?,?,?,?)',(dr['service_date'],rid,dr['driver_id'],to_driver,start_id,reason,note,u['id'],htime)); hid=cur.lastrowid
    targets=c.execute("SELECT x.id FROM deliveries x JOIN branches b ON b.id=x.branch_id WHERE x.daily_route_id=? AND b.stop_order>=? AND x.status!='STOP_COMPLETED' ORDER BY b.stop_order",(rid,start['stop_order'])).fetchall()
    for x in targets: c.execute('INSERT INTO delivery_driver_assignments(delivery_id,driver_id,assigned_at,handoff_id) VALUES(?,?,?,?) ON CONFLICT(delivery_id) DO UPDATE SET driver_id=excluded.driver_id,assigned_at=excluded.assigned_at,handoff_id=excluded.handoff_id',(x['id'],to_driver,htime,hid))
    audit(c,'USER',u['id'],u['role'],'DRIVER_HANDOFF','DAILY_ROUTE',rid,after={'from_driver_id':dr['driver_id'],'to_driver_id':to_driver,'start_delivery_id':start_id,'count':len(targets),'reason':reason,'note':note});c.commit();c.close();await publish({'type':'route.handoff','id':rid});return {'ok':True,'reassigned':len(targets)}

@app.get('/api/driver/handoff-notices')
def driver_handoff_notices(req:Request):
    s=driver_auth(req);c=db();rows=[dict(x) for x in c.execute("SELECT h.*,r.code route_code,fd.name from_driver,b.name start_branch,(SELECT COUNT(*) FROM delivery_driver_assignments a JOIN deliveries x2 ON x2.id=a.delivery_id WHERE a.handoff_id=h.id AND x2.status!='STOP_COMPLETED') remaining FROM route_handoffs h JOIN daily_routes dr ON dr.id=h.daily_route_id JOIN routes r ON r.id=dr.route_id JOIN drivers fd ON fd.id=h.from_driver_id JOIN deliveries x ON x.id=h.start_delivery_id JOIN branches b ON b.id=x.branch_id WHERE h.service_date=? AND h.to_driver_id=? ORDER BY h.id DESC",(today(),s['driver_id'])).fetchall()];c.close();return rows

@app.get('/api/daily-routes/{rid}/handoff-options')
def handoff_options(rid:int,req:Request):
    require_user(req,['ADMIN']);c=db();stops=[dict(x) for x in c.execute("SELECT x.id,b.code,b.name,b.stop_order,x.status FROM deliveries x JOIN branches b ON b.id=x.branch_id WHERE x.daily_route_id=? AND x.status!='STOP_COMPLETED' ORDER BY b.stop_order",(rid,)).fetchall()];drivers=[dict(x) for x in c.execute('SELECT id,name FROM drivers WHERE active=1 ORDER BY name').fetchall()];c.close();return {'stops':stops,'drivers':drivers}

@app.post('/api/driver/routes/{rid}/segment-sign')
async def segment_sign(rid:int,req:Request):
    s=driver_auth(req);p=await req.json();sig=p.get('signature') or ''
    if not sig: raise HTTPException(400,'請先簽名')
    c=db(); assigned=c.execute("SELECT COUNT(*) n,SUM(CASE WHEN x.status='STOP_COMPLETED' THEN 1 ELSE 0 END) done FROM deliveries x JOIN daily_routes dr ON dr.id=x.daily_route_id LEFT JOIN delivery_driver_assignments a ON a.delivery_id=x.id WHERE dr.id=? AND COALESCE(a.driver_id,dr.driver_id)=?",(rid,s['driver_id'])).fetchone()
    if not assigned or not assigned['n'] or assigned['done']<assigned['n']: c.close(); raise HTTPException(409,'您的配送區段尚未全部完成')
    c.execute('INSERT INTO route_segment_signatures(service_date,daily_route_id,driver_id,signature,signed_at) VALUES(?,?,?,?,?) ON CONFLICT(service_date,daily_route_id,driver_id) DO UPDATE SET signature=excluded.signature,signed_at=excluded.signed_at',(today(),rid,s['driver_id'],sig,now()));audit(c,'DRIVER',s['driver_id'],'DRIVER','DRIVER_SEGMENT_SIGN','DAILY_ROUTE',rid);c.commit();c.close();return {'ok':True}

@app.post('/api/email/smtp-test')
async def smtp_test(req:Request):
    require_user(req,['ADMIN']);p=await req.json();email=(p.get('email') or '').strip()
    if '@' not in email: raise HTTPException(400,'Email格式錯誤')
    try:smtp_send('圖書物流系統 SMTP 測試','lib.moving-match.com SMTP 郵件服務測試成功。',[email])
    except Exception as e:raise HTTPException(500,str(e))
    return {'ok':True}
