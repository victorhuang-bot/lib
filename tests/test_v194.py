"""Offline API/report regression tests. All DBs are newly created under /tmp.
Never import app.main against inherited DATABASE_URL or the tracked data/app.db.
"""
import base64
import io
import os
from pathlib import Path
import shutil
import socket
import sqlite3
import tempfile
import unittest
import csv
from contextlib import contextmanager
from concurrent.futures import ThreadPoolExecutor
from unittest.mock import patch

from fastapi.testclient import TestClient
from openpyxl import load_workbook

_bootstrap = tempfile.TemporaryDirectory(prefix='movingmatch-v194-bootstrap-')
with patch.dict(os.environ, {'DATABASE_URL': '', 'DATA_DIR': _bootstrap.name, 'APP_ENV': 'development'}, clear=True):
    with patch.object(socket.socket, 'connect', side_effect=AssertionError('Tests prohibit outbound network')):
        from app import main
        main.init_db()
        main.seed_demo_data()

LABELS = ['公文－總館送出', '公文－分館送出', '圖書－倉庫送出', '圖書－分館送出']
PNG = 'data:image/png;base64,' + base64.b64encode(b'not-used-by-report').decode()

class V194Tests(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory(prefix='movingmatch-v194-test-')
        self.addCleanup(self.tmp.cleanup)
        self.path = str(Path(self.tmp.name) / 'test.db')
        source = sqlite3.connect(str(Path(_bootstrap.name) / 'app.db'))
        target = sqlite3.connect(self.path)
        source.backup(target); source.close(); target.close()
        def connect():
            con = sqlite3.connect(self.path, timeout=5)
            con.row_factory = sqlite3.Row
            con.execute('PRAGMA foreign_keys=ON')
            return con
        self.connect = connect
        self.db_patch = patch.object(main, 'db', side_effect=connect)
        self.db_patch.start(); self.addCleanup(self.db_patch.stop)
        self.network = patch.object(socket.socket, 'connect', side_effect=AssertionError('Tests prohibit outbound network'))
        self.network.start(); self.addCleanup(self.network.stop)
        self.client = TestClient(main.app)
        self.addCleanup(self.client.close)
        with self.connection() as c:
            c.execute('UPDATE branches SET active=1 WHERE id IN (1,2)')
            for bid, token in [(1, 'branch-one'), (2, 'branch-two')]:
                c.execute('INSERT INTO branch_sessions(token_hash,branch_id,delivery_id,expires_at) VALUES(?,?,NULL,?)', (main.thash(token), bid, main.future_iso(hours=1)))
            c.execute("INSERT INTO users(username,password_hash,role) VALUES('test-admin','unused','ADMIN')")
            uid=c.execute("SELECT id FROM users WHERE username='test-admin'").fetchone()['id']
            c.execute('INSERT INTO sessions(token_hash,user_id,expires_at) VALUES(?,?,?)',(main.thash('admin'),uid,main.future_iso(hours=1)))
        self.headers={'Authorization':'Bearer branch-one'}
        self.other={'Authorization':'Bearer branch-two'}
        self.admin={'Authorization':'Bearer admin'}

    @contextmanager
    def connection(self):
        con=self.connect()
        try:
            with con: yield con
        finally: con.close()

    def create(self, name='王小明', headers=None):
        return self.client.post('/api/branch-session/signers', json={'name':name}, headers=headers or self.headers)

    def test_branch_scoped_add_rename_deactivate_and_audit(self):
        first=self.create(); self.assertEqual(first.status_code,200,first.text)
        sid=first.json()['id']; self.assertIsInstance(sid,int)
        self.assertEqual(self.client.get('/api/branch-session/signers',headers=self.other).json(),[])
        second=self.create(headers=self.other); self.assertEqual(second.status_code,200)
        result=self.client.patch(f'/api/branch-session/signers/{sid}',json={'name':'陳小華'},headers=self.headers)
        self.assertEqual(result.status_code,200,result.text)
        self.assertEqual(result.json()['name'],'陳小華')
        self.assertEqual(self.client.patch(f'/api/branch-session/signers/{sid}',json={'active':False},headers=self.headers).status_code,200)
        self.assertEqual(self.client.get('/api/branch-session/signers',headers=self.headers).json(),[])
        rows=self.client.get('/api/branch-session/signers?include_inactive=true',headers=self.headers).json()
        self.assertEqual([(r['name'],r['active']) for r in rows],[('陳小華',0)])
        self.assertEqual(self.client.get('/api/branch-session/signers',headers=self.other).json()[0]['name'],'王小明')
        with self.connection() as c:
            actions=[r['action'] for r in c.execute("SELECT action FROM audit_logs WHERE entity_type='BRANCH_SIGNER' AND actor_id='1' ORDER BY id")]
        self.assertEqual(actions,['ADD_BRANCH_SIGNER','UPDATE_BRANCH_SIGNER','UPDATE_BRANCH_SIGNER'])

    def test_cross_branch_and_other_roles_cannot_modify(self):
        sid=self.create().json()['id']
        for payload in [{'name':'竄改'}, {'active':False}]:
            self.assertEqual(self.client.patch(f'/api/branch-session/signers/{sid}',json=payload,headers=self.other).status_code,404)
        for headers in [{},self.admin,{'Authorization':'Bearer invalid'}]:
            for method,path,payload in [('GET','/api/branch-session/signers',None),('POST','/api/branch-session/signers',{'name':'不允許'}),('PATCH',f'/api/branch-session/signers/{sid}',{'name':'不允許'})]:
                self.assertEqual(self.client.request(method,path,json=payload,headers=headers).status_code,401)
        self.assertEqual(self.client.post('/api/branch-session/signers',json={'name':'覆寫','branch_id':2},headers=self.headers).status_code,400)

    def test_validation_duplicates_and_disabled_names(self):
        for invalid in ['', '  ', None, 123, ['name'], '字'*81]:
            self.assertEqual(self.create(invalid).status_code,400)
        sid=self.create('  王小明  ').json()['id']
        self.assertEqual(self.create('王小明').status_code,409)
        second=self.create('陳小華').json()['id']
        self.assertEqual(self.client.patch(f'/api/branch-session/signers/{second}',json={'name':'王小明'},headers=self.headers).status_code,409)
        self.assertEqual(self.client.patch(f'/api/branch-session/signers/{sid}',json={'active':False},headers=self.headers).status_code,200)
        self.assertEqual(self.create('王小明').status_code,409)
        for invalid in [{}, {'active':1}, {'name':' '}, {'branch_id':2}]:
            self.assertEqual(self.client.patch(f'/api/branch-session/signers/{second}',json=invalid,headers=self.headers).status_code,400)
        with self.connection() as c:
            self.assertEqual(c.execute('SELECT count(*) FROM branch_signers WHERE branch_id=1').fetchone()[0],2)

    def test_inactive_branch_and_expired_session_are_rejected(self):
        with self.connection() as c: c.execute('UPDATE branches SET active=0 WHERE id=1')
        self.assertEqual(self.create().status_code,403)
        with self.connection() as c: c.execute("UPDATE branch_sessions SET expires_at='2000-01-01' WHERE branch_id=2")
        self.assertEqual(self.create(headers=self.other).status_code,401)

    def test_roster_updates_preserve_all_delivery_history_and_counts(self):
        with self.connection() as c:
            c.execute("INSERT INTO deliveries(id,service_date,branch_id,status,document_final,document_return_final,outbound_final,inbound_final,signer_name,branch_signature,correction_signer_name,correction_signature) VALUES(990001,'2025-01-01',1,'STOP_COMPLETED',11,22,33,44,'王小明','old-signature','王小明','old-correction')")
            before=[tuple(r) for r in c.execute('SELECT * FROM deliveries ORDER BY id')]
        sid=self.create().json()['id']
        self.client.patch(f'/api/branch-session/signers/{sid}',json={'name':'新姓名','active':False},headers=self.headers)
        with self.connection() as c: after=[tuple(r) for r in c.execute('SELECT * FROM deliveries ORDER BY id')]
        self.assertEqual(before,after)

    def test_no_delivery_desktop_and_qr_login_still_allow_roster_management(self):
        with self.connection() as c:
            c.execute("UPDATE branches SET pin_hash=?,access_token_hash=? WHERE id=1",(main.hash_secret('2468'),main.thash('test-qr')))
            c.execute('DELETE FROM deliveries WHERE branch_id=1')
            code=c.execute('SELECT code FROM branches WHERE id=1').fetchone()['code']
        for endpoint,payload in [('desktop-login',{'branch_code':code,'pin':'2468'}),('verify',{'token':'test-qr','pin':'2468'})]:
            login=self.client.post('/api/branch-session/'+endpoint,json=payload)
            self.assertEqual(login.status_code,200,login.text)
            headers={'Authorization':'Bearer '+login.json()['session']}
            response=self.client.get('/api/branch-session/today',headers=headers)
            self.assertEqual(response.json()['status'],'NO_DELIVERY')
            self.assertEqual(self.client.get('/api/branch-session/tasks',headers=headers).status_code,200)
            self.assertEqual(self.client.get('/api/branch-session/signers',headers=headers).status_code,200)
            self.assertEqual(self.client.post('/api/branch-session/correct',json={},headers=headers).status_code,409)

    def report_fixture(self):
        with self.connection() as c:
            c.execute("INSERT INTO daily_routes(id,service_date,route_id,driver_id) VALUES(990001,'2025-02-01',1,NULL)")
            c.execute("INSERT INTO deliveries(id,service_date,daily_route_id,branch_id,status,document_final,document_return_final,outbound_final,inbound_final,signer_name,branch_signed_at) VALUES(990002,'2025-02-01',990001,1,'STOP_COMPLETED',11,22,33,44,'歷史姓名','2025-02-01T10:00:00+08:00')")

    def test_daily_monthly_excel_labels_map_to_original_quantities(self):
        self.report_fixture()
        for endpoint in ['/api/reports/daily.xlsx?service_date=2025-02-01','/api/reports/monthly.xlsx?month=2025-02']:
            result=self.client.get(endpoint,headers=self.admin)
            self.assertEqual(result.status_code,200,result.text[:80] if result.status_code!=200 else '')
            sheet=load_workbook(io.BytesIO(result.content)).active
            headers=[cell.value for cell in sheet[2]]
            values=[cell.value for cell in sheet[3]]
            for label,value in zip(LABELS,[11,22,33,44]): self.assertEqual(values[headers.index(label)],value)
        with self.connection() as c:
            legacy=main.report_rows(c,'2025-02-01','2025-02-01')[0]
            self.assertEqual([legacy[key] for key in ['公文送出','公文收回','圖書送出','圖書收回']],[11,22,33,44])
            payload=main.official_daily_payload(c,'2025-02-01')
        sheet=load_workbook(io.BytesIO(main.official_xlsx(payload,'2025-02-01'))).active
        self.assertEqual([cell.value for cell in sheet[2]][3:7],LABELS)
        self.assertEqual([cell.value for cell in sheet[3]][3:7],[11,22,33,44])
        self.assertEqual(payload['grand'],{'document':11,'document_return':22,'outbound':33,'inbound':44})

    def test_pdf_and_frontend_labels(self):
        self.report_fixture()
        with patch.object(main.canvas,'Canvas') as canvas:
            with self.connection() as c: main.pdf_bytes(main.report_rows(c,'2025-02-01','2025-02-01'),'測試')
            text=' '.join(str(call.args[-1]) for call in canvas.return_value.drawString.call_args_list)
            for label in LABELS: self.assertIn(label,text)
        for filename in ['index.html','branch.html','driver.html']:
            source=(main.STATIC/filename).read_text()
            for old in ['公文送出','公文收回','公文送回','圖書送出','圖書收回','圖書送回']: self.assertNotIn(old,source)
            for label in LABELS: self.assertIn(label,source)

    def test_concurrent_duplicate_creates_have_one_winner(self):
        with ThreadPoolExecutor(max_workers=2) as pool:
            results=list(pool.map(lambda _: self.create('同一姓名').status_code, range(2)))
        self.assertEqual(sorted(results),[200,409])
        self.assertEqual(len(self.client.get('/api/branch-session/signers',headers=self.headers).json()),1)

    def test_additive_migration_is_idempotent_and_preserves_data(self):
        sid=self.create().json()['id']
        with self.connection() as c:
            before=[tuple(r) for r in c.execute('SELECT * FROM branch_signers')]
            main.migrate_branch_signers(c)
            main.migrate_branch_signers(c)
            self.assertEqual(before,[tuple(r) for r in c.execute('SELECT * FROM branch_signers')])
        self.assertIsInstance(sid,int)

    def test_postgres_identity_adapter_returns_new_signer_id(self):
        from unittest.mock import Mock
        connection=Mock()
        cursor=connection.cursor.return_value
        cursor.fetchone.return_value={'id':71}
        pool=Mock(); pool.getconn.return_value=connection
        with patch.object(main,'pg_pool',return_value=pool):
            c=main._PgCompat('unused-test-url')
            result=c.execute('INSERT INTO branch_signers(branch_id,name,active,created_at,updated_at) VALUES(?,?,1,?,?)',(1,'測試','now','now'))
            self.assertEqual(result.lastrowid,71)
            sql=cursor.execute.call_args.args[0]
            self.assertIn('RETURNING id',sql)
            self.assertNotIn('?',sql)
            c.close()

    def test_csv_and_official_pdf_keep_all_four_quantities(self):
        self.report_fixture()
        with self.connection() as c:
            driver=c.execute('SELECT id FROM drivers LIMIT 1').fetchone()['id']
            c.execute('UPDATE daily_routes SET driver_id=? WHERE id=990001',(driver,))
        with patch.object(main,'today',return_value='2025-02-01'):
            response=self.client.get('/api/reports/today.csv',headers=self.admin)
        self.assertEqual(response.status_code,200)
        rows=list(csv.reader(io.StringIO(response.content.decode('utf-8-sig'))))
        for label,value in zip(LABELS,['11','22','33','44']): self.assertEqual(rows[1][rows[0].index(label)],value)
        with self.connection() as c: payload=main.official_daily_payload(c,'2025-02-01')
        with patch.object(main.canvas,'Canvas') as canvas:
            main.official_pdf(payload,'2025-02-01')
            text=' '.join(str(call.args[-1]) for call in canvas.return_value.drawString.call_args_list)
            for label in LABELS: self.assertIn(label,text)

if __name__=='__main__': unittest.main()
